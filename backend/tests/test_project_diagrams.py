"""도식 의미 입력의 프로젝트 왕복과 공통 렌더 연결. 합성 자료만 사용한다."""

from copy import deepcopy
import asyncio
import json
from pathlib import Path

import pytest
from pptx import Presentation

from slidecaptain.export.exporter import export_deck_data
from slidecaptain.layout.engine import build_render_plan
from slidecaptain.metrics.font_metrics import FontMetrics
from slidecaptain.models.deck import Deck
from slidecaptain.models.preset import Preset
from slidecaptain.pipeline.quality import assess_quality, QualityExportBlocked
from slidecaptain.pipeline.story import story_fingerprint


@pytest.fixture
def project_input():
    sample = json.loads((Path(__file__).parent / 'fixtures/q3a-diagram.json').read_text('utf-8'))
    diagram = sample['diagram']
    chapter_id = diagram['id']
    deck = {
        'schema_version': 1,
        'meta': {'title': '합성 도식 보고'},
        'structure': {
            'chapters': [
                {'id': 'cover', 'topic': '합성 보고', 'template': 'cover'},
                {'id': chapter_id, 'topic': '요청 처리 흐름', 'template': 'diagram'},
            ],
            'story_plan': {
                'brief': {'decision_question': '요청을 어떻게 처리하는가?'},
                'evidence': sample['evidence'],
                'claims': [{'id': 'claim', 'statement': '접수 뒤 검토한다.', 'kind': 'fact',
                            'evidence_ids': ['e1', 'e2', 'e3'], 'caveats': []}],
                'answer_claim_ids': ['claim'],
                'chapters': [{'chapter_id': chapter_id, 'role': 'answer', 'claim_ids': ['claim']}],
                'unanswered_questions': [],
                'input_fingerprint': '0' * 64,
            },
        },
        # 저장 배열 순서와 렌더 순서를 의도적으로 다르게 둔다.
        'slides': [
            {'chapter_id': chapter_id, 'eyebrow': '합성 운영', 'subtitle': '접수부터 검토까지',
             'slots': {'template': 'diagram', 'diagram': diagram, 'footnote': '합성 자료'}},
            {'chapter_id': 'cover', 'slots': {'template': 'cover', 'title': '합성 보고'}},
        ],
    }
    return {'deck': deck, 'sources': sample['sources']}


def validated(sample):
    deck = Deck.model_validate(sample['deck'])
    plan = deck.structure.story_plan
    plan.input_fingerprint = story_fingerprint(plan, deck.meta, deck.structure.chapters, sample['sources'])
    return deck


def render(deck, preset=None):
    return build_render_plan(deck, preset or Preset(), FontMetrics.load_default())


def test_project_roundtrip_snapshot_and_etag_preserve_semantics(project_input, store):
    deck = validated(project_input)
    store.create_project('diagram')
    initial = store.deck_etag('diagram')
    tag = store.save_deck('diagram', deck, expected_etag=initial)
    saved, loaded_tag = store.load_deck_with_etag('diagram')
    assert tag == loaded_tag and tag != initial
    assert saved.model_dump() == deck.model_dump()
    store.snapshot_now('diagram')
    snapshot = store.list_snapshots('diagram')[-1].id
    changed = deepcopy(deck)
    changed.slides[0].slots.diagram.nodes[0].content = '수정 접수'
    store.save_deck('diagram', changed, expected_etag=tag)
    restored, _ = store.restore_snapshot('diagram', snapshot)
    assert restored == deck
    assert 'render_plan' not in restored.model_dump_json()
    assert 'input_fingerprint' not in restored.slides[0].slots.model_dump()


@pytest.mark.parametrize('mutation', ['missing_evidence', 'missing_ledger', 'wrong_id', 'render', 'review', 'coordinates'])
def test_invalid_diagram_payload_is_rejected_without_saving(project_input, client, store, mutation):
    store.create_project('diagram')
    before = store.deck_etag('diagram')
    payload = project_input['deck']
    slots = payload['slides'][0]['slots']
    if mutation == 'missing_evidence':
        slots['diagram']['edges'][0]['evidence_ids'] = ['absent']
    elif mutation == 'missing_ledger':
        payload['structure']['story_plan'] = None
    elif mutation == 'wrong_id':
        slots['diagram']['id'] = 'other'
    elif mutation == 'render':
        slots['render_plan'] = {'slides': []}
    elif mutation == 'review':
        slots['diagram']['review'] = {'semantic': 'passed'}
    else:
        slots['diagram']['nodes'][0]['x'] = 100
    response = client.put('/api/projects/diagram/deck', json=payload)
    assert response.status_code == 422
    assert store.deck_etag('diagram') == before


def test_project_uses_structure_order_page_numbers_and_shared_render(project_input, client, store):
    store.create_project('diagram')
    deck = validated(project_input)
    assert client.put('/api/projects/diagram/deck', json=deck.model_dump(mode='json')).status_code == 200
    reopened = client.get('/api/projects/diagram/deck')
    measured = client.post('/api/render-plan', json=reopened.json())
    fetched = client.get('/api/projects/diagram/render-plan')
    assert measured.status_code == fetched.status_code == 200
    assert measured.json() == fetched.json() == render(deck).model_dump(mode='json')
    slides = fetched.json()['slides']
    assert [s['chapter_id'] for s in slides] == ['cover', 'synthetic-flow']
    page = slides[1]['diagram']
    number = next(h for h in page['headers'] if h['name'].endswith(':page_number'))
    assert number['paras'][0]['text'] == '2'
    assert page['review'] == dict.fromkeys(['semantic', 'browser', 'powerpoint', 'reader'], 'not_run')
    assert [n['node'] for n in page['layout']['nodes']] == project_input['deck']['slides'][0]['slots']['diagram']['nodes']


@pytest.mark.parametrize('changed', ['node', 'evidence', 'title', 'eyebrow', 'subtitle', 'footnote', 'order', 'preset'])
def test_semantic_and_display_edits_recompute_fingerprints(project_input, changed):
    deck = validated(project_input)
    preset = Preset()
    old_render = render(deck, preset)
    old = old_render.slides[1].diagram.input_fingerprint
    old_quality = assess_quality(deck, preset, old_render, sources=project_input['sources']).input_fingerprint
    if changed == 'node':
        deck.slides[0].slots.diagram.nodes[0].content = '신규 요청 접수'
    elif changed == 'evidence':
        deck.structure.story_plan.evidence[0].excerpt = '근거 수정'
    elif changed == 'title':
        deck.structure.chapters[1].topic = '접수 및 검토'
    elif changed in ('eyebrow', 'subtitle'):
        setattr(deck.slides[0], changed, '수정 문구')
    elif changed == 'footnote':
        deck.slides[0].slots.footnote = '수정 자료'
    elif changed == 'order':
        deck.structure.chapters.reverse()
    else:
        preset.colors.accent1 = '112233'
    fresh = render(deck, preset)
    page = next(s.diagram for s in fresh.slides if s.diagram is not None)
    assert page.input_fingerprint != old
    assert assess_quality(deck, preset, fresh, sources=project_input['sources']).input_fingerprint != old_quality
    assert set(page.review.model_dump().values()) == {'not_run'}


def test_blocked_diagram_has_422_and_no_export_files(project_input, client, store):
    store.create_project('diagram')
    payload = project_input['deck']
    payload['slides'][0]['slots']['diagram']['edges'][0]['to_node_id'] = 'pilot'
    assert client.put('/api/projects/diagram/deck', json=payload).status_code == 200
    for method, path, body in (
        ('get', '/api/projects/diagram/render-plan', None),
        ('post', '/api/render-plan', payload),
        ('post', '/api/projects/diagram/export', None),
    ):
        response = getattr(client, method)(path, **({'json': body} if body else {}))
        assert response.status_code == 422
        assert '도식' in response.json()['detail']
    assert list(store.exports_dir('diagram').iterdir()) == []


def test_mixed_deck_exports_editable_diagram_but_not_final(project_input, tmp_path):
    deck = validated(project_input)
    path = export_deck_data(deck, tmp_path / 'exports', sources=project_input['sources'])
    pptx = Presentation(path)
    assert len(pptx.slides) == 2
    assert any(shape.name == 'synthetic-flow:node:intake' for shape in pptx.slides[1].shapes)
    quality = json.loads(path.with_suffix('.quality.json').read_text('utf-8'))
    assert quality['final_export_allowed'] is False
    with pytest.raises(QualityExportBlocked):
        export_deck_data(deck, tmp_path / 'final', final=True, sources=project_input['sources'])
    assert not (tmp_path / 'final').exists()


@pytest.mark.parametrize('route', ['generate/chapter/synthetic-flow', 'generate/chapter/synthetic-flow/condense',
                                 'generate/chapter/cover/condense'])
def test_diagram_generation_rejected_before_provider_setup(project_input, client, store, route):
    store.create_project('diagram')
    store.save_deck('diagram', validated(project_input))
    response = client.post(f'/api/projects/diagram/{route}',
                           json={'slots': project_input['deck']['slides'][0]['slots']},
                           headers={'X-AI-Consent': 'SlideCaptain'})
    assert response.status_code == 422
    assert '도식' in response.json()['detail']


def test_diagram_chapter_cannot_be_silently_skipped(project_input, client, store):
    store.create_project('diagram')
    payload = project_input['deck']
    payload['slides'] = payload['slides'][1:]
    response = client.put('/api/projects/diagram/deck', json=payload)
    assert response.status_code == 422
    assert '도식' in response.json()['detail']


def test_whole_structure_generation_cannot_replace_diagram_ledger(project_input, client, store):
    store.create_project('diagram')
    store.save_deck('diagram', validated(project_input))
    response = client.post('/api/projects/diagram/generate/structure', json={},
                           headers={'X-AI-Consent': 'SlideCaptain'})
    assert response.status_code == 422
    assert '도식' in response.json()['detail']


@pytest.mark.parametrize('boundary', ['render', 'save'])
def test_in_memory_mutation_is_revalidated_without_losing_extra_fields(project_input, store, boundary):
    deck = validated(project_input)
    deck.slides[0].slots.diagram.nodes[0].__dict__['x'] = 100
    if boundary == 'render':
        with pytest.raises(ValueError):
            render(deck)
    else:
        store.create_project('diagram')
        before = store.deck_etag('diagram')
        with pytest.raises(ValueError):
            store.save_deck('diagram', deck)
        assert store.deck_etag('diagram') == before
        assert store.list_snapshots('diagram') == []


def test_proposals_can_use_empty_ledger_but_title_is_required(project_input):
    payload = project_input['deck']
    payload['structure']['story_plan'] = None
    diagram = payload['slides'][0]['slots']['diagram']
    for node in diagram['nodes']:
        node.update(kind='proposal', evidence_ids=[])
    for edge in diagram['edges']:
        edge.update(relation='proposal', evidence_ids=[])
    deck = Deck.model_validate(payload)
    assert render(deck).slides[1].diagram is not None
    payload['structure']['chapters'][1]['topic'] = '  '
    with pytest.raises(ValueError, match='제목'):
        Deck.model_validate(payload)


def test_shared_project_fixture_and_visible_numeric_fields():
    from slidecaptain.pipeline.numeric_review import assess_numeric_review

    sample = json.loads((Path(__file__).parent / 'fixtures/q3b-project.json').read_text('utf-8'))
    deck = Deck.model_validate(sample['deck'])
    assert render(deck).model_dump(mode='json') == sample['render_plan']
    assert assess_numeric_review(deck, sample['sources']).numeric_fields == 0
    diagram = deck.slides[0].slots.diagram
    diagram.nodes[0].content = '요청 10건 접수'
    diagram.nodes[0].caveats = ['처리 시간 2일은 미확인']
    diagram.edges[0].label = '검토 1회'
    deck.slides[0].slots.footnote = '합성 자료 3종'
    assert assess_numeric_review(deck, sample['sources']).numeric_fields == 4


def test_direct_generation_service_never_calls_provider_for_diagrams(project_input):
    from unittest.mock import AsyncMock
    from slidecaptain.pipeline.service import GenerationService, DiagramGenerationUnsupported

    provider = AsyncMock()
    svc = GenerationService(provider, FontMetrics.load_default())
    deck = validated(project_input)
    for method, args in (
        (svc.generate_chapter, (deck, 'synthetic-flow', {}, Preset())),
        (svc.condense_chapter, (deck, 'synthetic-flow', deck.slides[0].slots, {}, Preset())),
    ):
        with pytest.raises(DiagramGenerationUnsupported):
            asyncio.run(method(*args))
    provider.complete.assert_not_called()


@pytest.mark.parametrize('planned', [False, True])
def test_structure_generation_schema_and_parser_do_not_advertise_diagrams(planned):
    from slidecaptain.models.deck import DeckMeta
    from slidecaptain.models.story import ReportBrief
    from slidecaptain.pipeline.prompts import structure_response_schema
    from slidecaptain.pipeline.service import GenerationService
    from test_story_plan import BRIEF, PAYLOAD, SOURCES, StubProvider

    assert '"diagram"' not in json.dumps(structure_response_schema(planned=planned))
    if planned:
        payload = deepcopy(PAYLOAD)
        payload['chapters'][0]['template'] = 'diagram'
    else:
        payload = {'chapters': [{'topic': '도식', 'conclusion': '', 'template': 'diagram', 'source_refs': []}]}
    provider = StubProvider([payload, payload])
    svc = GenerationService(provider, FontMetrics.load_default())
    result = asyncio.run(svc.generate_structure(DeckMeta(title='합성 보고'), SOURCES,
                                               brief=ReportBrief(**BRIEF) if planned else None))
    assert result.status == 'format_error'
    assert result.structure is None
