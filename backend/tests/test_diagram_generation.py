"""AI 도식 후보의 요청 경계, 보호 저장본, 사용량과 수동 적용 왕복. 외부 호출 없음."""

import asyncio
from copy import deepcopy
import hashlib
import json
from pathlib import Path

from fastapi.testclient import TestClient
from pptx import Presentation
import pytest

from slidecaptain.metrics.font_metrics import FontMetrics
from slidecaptain.models.deck import Deck
from slidecaptain.pipeline.auth_status import LoginStatus
from slidecaptain.pipeline.connections import AIConnections, AISelection, ConnectionConflict, LoginAttempt, ModelOption
from slidecaptain.pipeline.provider import CallUsage, ProviderCallFailed, ProviderResponse
from slidecaptain.pipeline.rewrite import sources_fingerprint
from slidecaptain.pipeline.story import story_fingerprint
from slidecaptain.server.app import create_app


ROUTE = "/api/projects/diagram/generate/diagram"
REQUEST = {
    "chapter_id": "new-diagram", "topic": "요청 처리 초안", "role": "evidence",
    "claim_ids": ["claim"], "instructions": "원문의 관계와 조건을 보존",
}


class StubProvider:
    model = "synthetic-model"

    def __init__(self, payload, *, responses=None, effect=None):
        self.payload = payload
        self.responses = iter(responses) if responses is not None else None
        self.effect = effect
        self.calls = []

    async def complete(self, prompt, schema):
        self.calls.append((prompt, schema))
        if self.effect:
            self.effect()
        value = next(self.responses) if self.responses is not None else ProviderResponse(
            structured=deepcopy(self.payload), raw_text="합성 도식 후보",
        )
        if isinstance(value, Exception):
            raise value
        return value


@pytest.fixture
def diagram_input(store):
    sample = json.loads((Path(__file__).parent / "fixtures/q3b-project.json").read_text("utf-8"))
    sources = sample["sources"]
    source_id = next(iter(sources))
    sources[source_id] += "\n발췌 밖 내용 FULL_SOURCE_NOT_SENT 777"
    sources["unselected.md"] = "선택하지 않은 자료 UNSELECTED_SOURCE_NOT_SENT 999"
    plan = sample["deck"]["structure"]["story_plan"]
    for item in plan["evidence"]:
        item["source_revision"] = hashlib.sha256(sources[item["source_id"]].encode()).hexdigest()
    plan["evidence"].append({
        **deepcopy(plan["evidence"][0]), "id": "unselected-evidence", "source_id": "unselected.md",
        "source_revision": hashlib.sha256(sources["unselected.md"].encode()).hexdigest(),
        "excerpt": sources["unselected.md"],
    })
    plan["claims"].append({
        "id": "unselected-claim", "statement": "UNSELECTED_CLAIM_NOT_SENT", "kind": "fact",
        "evidence_ids": ["unselected-evidence"], "caveats": [],
    })
    plan["chapters"][0]["claim_ids"].append("unselected-claim")
    deck = Deck.model_validate(sample["deck"])
    refresh_plan(deck, sources)
    store.create_project("diagram")
    for filename, text in sources.items():
        store.write_source("diagram", filename, text)
    store.save_deck("diagram", deck, snapshot=False)
    diagram = deck.slides[0].slots.diagram.model_dump(mode="json")
    return deck, sources, {key: diagram[key] for key in ("nodes", "edges")}


def refresh_plan(deck, sources):
    plan = deck.structure.story_plan
    plan.input_fingerprint = story_fingerprint(plan, deck.meta, deck.structure.chapters, sources)


def app_client(store, provider=None, manager=None):
    return TestClient(create_app(store, provider=provider, ai_connections=manager,
                                 login_checker=lambda: LoginStatus(logged_in=True)), headers={
        "X-Requested-With": "SlideCaptain", "X-AI-Consent": "SlideCaptain",
        "If-Match": f'"{store.deck_etag("diagram")}"',
    })


def usage_records(store):
    path = store.root / "diagram" / "ai-usage.jsonl"
    return [json.loads(line) for line in path.read_text("utf-8").splitlines()] if path.exists() else []


def usage(**updates):
    payload = {name: None for name in CallUsage.model_fields}
    payload.update(model="synthetic-model", token_source="usage", input_tokens=10, output_tokens=5,
                   cache_read_tokens=0, cache_creation_tokens=0)
    payload.update(updates)
    return CallUsage(**payload)


def test_preview_returns_server_owned_spec_without_saving_or_sending_other_context(diagram_input, store):
    deck, sources, payload = diagram_input
    provider = StubProvider(payload)
    etag = store.deck_etag("diagram")
    client = app_client(store, provider)
    response = client.post(ROUTE, json=REQUEST)
    assert response.status_code == 200, response.text
    result = response.json()
    assert result["status"] == "ok"
    assert result["diagram"] == {
        "version": "q3a-v1", "id": REQUEST["chapter_id"],
        "canvas": {"page_size": "preset", "reading_profile": "report"},
        **payload, "layout_variant": "flow_horizontal",
    }
    assert result["base_etag"] == f'"{etag}"'
    assert result["sources_fingerprint"] == sources_fingerprint(sources)
    assert result["format_retried"] is False
    assert result["usage"]["calls"] == 1
    assert store.load_deck("diagram") == deck
    assert store.deck_etag("diagram") == etag
    assert store.list_snapshots("diagram") == []
    prompt, schema = provider.calls[0]
    for expected in (REQUEST["topic"], REQUEST["instructions"], deck.meta.title,
                     deck.structure.story_plan.brief.decision_question, "접수 뒤 검토한다.",
                     deck.structure.story_plan.evidence[2].excerpt):
        assert expected in prompt
    for omitted in ("FULL_SOURCE_NOT_SENT", "UNSELECTED_SOURCE_NOT_SENT", "UNSELECTED_CLAIM_NOT_SENT",
                    "synthetic-flow", "접수부터 검토까지"):
        assert omitted not in prompt
    assert "인과" in prompt and "추정" in prompt
    assert set(schema["properties"]) == {"nodes", "edges"}
    assert schema["additionalProperties"] is False
    record, = usage_records(store)
    assert record["kind"] == "diagram"
    assert record["chapter_id"] == REQUEST["chapter_id"]
    assert record["outcome"] == "ok"
    assert record["summary"] == result["usage"]


@pytest.mark.parametrize("field,value", [
    ("chapter_id", "synthetic-flow"), ("chapter_id", "cover"), ("chapter_id", "bad/id"),
    ("topic", "  "), ("topic", "제" * 501), ("role", "cover"), ("role", "divider"),
    ("claim_ids", []), ("claim_ids", ["claim", "claim"]), ("claim_ids", ["missing"]),
    ("claim_ids", [1]), ("instructions", "가" * 8001), ("evidence", []), ("deck", {}),
    ("nodes", []),
], ids=["existing-diagram", "existing-cover", "invalid-id", "blank-topic", "long-topic", "cover-role",
        "divider-role", "no-claims", "duplicate-claims", "unknown-claim", "invalid-claim", "long-instructions",
        "forged-evidence", "forged-deck", "forged-nodes"])
def test_bad_request_is_rejected_before_provider(diagram_input, store, field, value):
    deck, _, payload = diagram_input
    provider = StubProvider(payload)
    response = app_client(store, provider).post(ROUTE, json={**REQUEST, field: value})
    assert response.status_code == 422, response.text
    assert provider.calls == []
    assert usage_records(store) == []
    assert store.load_deck("diagram") == deck


@pytest.mark.parametrize("defect,status", [
    ("absent", 422), ("stale", 409), ("source_deleted", 409), ("source_changed", 409),
    ("source_added", 409), ("forged_hash", 409), ("forged_excerpt", 409), ("forged_locator", 409),
])
def test_plan_and_actual_source_must_be_current_before_provider(diagram_input, store, defect, status):
    deck, sources, payload = diagram_input
    if defect == "absent":
        # Existing diagrams require an evidence ledger, so use an ordinary legacy deck.
        deck = Deck(meta=deck.meta)
    elif defect == "stale":
        deck.structure.story_plan.input_fingerprint = "0" * 64
    elif defect == "source_deleted":
        (store.root / "diagram" / "sources" / next(iter(sources))).unlink()
    elif defect == "source_changed":
        store.write_source("diagram", next(iter(sources)), "새 원문")
    elif defect == "source_added":
        store.write_source("diagram", "added.md", "새 자료")
    else:
        item = deck.structure.story_plan.evidence[0]
        if defect == "forged_hash": item.source_revision = "0" * 64
        if defect == "forged_excerpt": item.excerpt = "현재 원문과 다른 발췌"
        if defect == "forged_locator": item.locator.line_end = 999
        refresh_plan(deck, sources)  # A freshly forged fingerprint cannot approve stale evidence.
    store.save_deck("diagram", deck, snapshot=False)
    provider = StubProvider(payload)
    response = app_client(store, provider).post(ROUTE, json=REQUEST)
    assert response.status_code == status, response.text
    assert provider.calls == []
    assert usage_records(store) == []


@pytest.mark.parametrize("header,status", [
    ("X-Requested-With", 403), ("X-AI-Consent", 428), ("If-Match", 428),
])
def test_required_gates_precede_generation(diagram_input, store, header, status):
    provider = StubProvider(diagram_input[2])
    client = app_client(store, provider)
    del client.headers[header]
    assert client.post(ROUTE, json=REQUEST).status_code == status
    assert provider.calls == []


def test_stale_etag_precedes_generation_setup(diagram_input, store):
    response = app_client(store).post(ROUTE, json=REQUEST, headers={"If-Match": '"old"'})
    assert response.status_code == 412, response.text
    assert usage_records(store) == []


@pytest.mark.parametrize("defect", [
    "not_selected", "unknown_evidence", "duplicate_node", "duplicate_edge", "unknown_node",
    "missing_relation_support", "missing_uncertainty", "coordinates", "fixed_id", "new_evidence",
    "role", "claim_ids", "cycle", "no_nodes", "no_edges",
])
def test_invalid_ai_output_retries_once_and_preserves_saved_deck(diagram_input, store, defect):
    deck, _, payload = diagram_input
    if defect == "not_selected": payload["nodes"][0]["evidence_ids"] = ["unselected-evidence"]
    if defect == "unknown_evidence": payload["nodes"][0]["evidence_ids"] = ["invented"]
    if defect == "duplicate_node": payload["nodes"][1]["id"] = payload["nodes"][0]["id"]
    if defect == "duplicate_edge": payload["edges"].append(deepcopy(payload["edges"][0]))
    if defect == "unknown_node": payload["edges"][0]["to_node_id"] = "absent"
    if defect == "missing_relation_support": payload["edges"][0]["evidence_ids"] = []
    if defect == "missing_uncertainty": payload["nodes"][0].update(kind="unknown", caveats=[])
    if defect == "coordinates": payload["nodes"][0]["x"] = 1
    if defect == "fixed_id": payload["id"] = "forged"
    if defect == "new_evidence": payload["evidence"] = []
    if defect == "role": payload["role"] = "answer"
    if defect == "claim_ids": payload["claim_ids"] = ["unselected-claim"]
    if defect == "cycle":
        payload["edges"].append({**payload["edges"][0], "id": "cycle", "from_node_id": "review", "to_node_id": "intake"})
    if defect == "no_nodes": payload["nodes"] = []
    if defect == "no_edges": payload["edges"] = []
    provider = StubProvider(payload)
    response = app_client(store, provider).post(ROUTE, json=REQUEST)
    assert response.status_code == 200, response.text
    result = response.json()
    assert result["status"] == "format_error"
    assert result["diagram"] is None
    assert result["raw_text"] == "합성 도식 후보"
    assert result["format_retried"] is True
    assert len(provider.calls) == 2
    assert result["usage"]["calls"] == 2
    assert [r["purpose"] for r in result["usage"]["records"]] == ["generate", "format_retry"]
    assert usage_records(store)[0]["outcome"] == "format_error"
    assert store.load_deck("diagram") == deck
    assert store.list_snapshots("diagram") == []


def test_retry_success_preserves_unknown_usage_and_does_not_normalize_text(diagram_input, store):
    _, _, payload = diagram_input
    payload["nodes"][0]["content"] = "요청  접수"
    provider = StubProvider(payload, responses=[
        ProviderResponse(structured={}, raw_text="bad", usage=usage()),
        ProviderResponse(structured=payload, raw_text="good", usage=usage(input_tokens=None, output_tokens=0)),
    ])
    response = app_client(store, provider).post(ROUTE, json=REQUEST)
    assert response.status_code == 200, response.text
    result = response.json()
    assert result["status"] == "ok" and result["format_retried"] is True
    assert result["diagram"]["nodes"][0]["content"] == "요청  접수"
    assert result["usage"]["input_tokens"] is None
    assert result["usage"]["output_tokens"] == 5
    assert "input_tokens" in result["usage"]["missing"]
    assert result["usage"]["cache_read_tokens"] == 0
    assert usage_records(store)[0]["summary"] == result["usage"]


@pytest.mark.parametrize("failed_attempt", [1, 2])
def test_provider_failure_records_usage_without_saving(diagram_input, store, failed_attempt):
    deck, _, payload = diagram_input
    responses = ([ProviderResponse(structured={}, raw_text="bad")] if failed_attempt == 2 else [])
    responses.append(ProviderCallFailed("합성 호출 실패", usage=usage()))
    provider = StubProvider(payload, responses=responses)
    response = app_client(store, provider).post(ROUTE, json=REQUEST)
    assert response.status_code == 503, response.text
    record, = usage_records(store)
    assert record["kind"] == "diagram" and record["outcome"] == "failed"
    assert record["summary"]["calls"] == failed_attempt
    assert record["summary"]["failed_calls"] == 1
    assert store.load_deck("diagram") == deck
    assert store.list_snapshots("diagram") == []


def test_numbers_only_use_selected_registered_excerpts(diagram_input, store):
    payload = diagram_input[2]
    payload["nodes"][0]["content"] = "접수 777건, 검토 999건"
    payload["nodes"][2]["caveats"] = ["비용 321원 미확인"]
    payload["edges"][1]["label"] = "알림 123회 제안"
    response = app_client(store, StubProvider(payload)).post(ROUTE, json=REQUEST)
    assert response.status_code == 200, response.text
    assert set(response.json()["unverified_numbers"]) == {"777", "999", "321", "123"}


@pytest.mark.parametrize("change", ["deck", "deck_invalid", "source", "source_too_large", "source_invalid",
                                     "add_source", "remove_source", "remove_all_sources"])
def test_changes_while_awaiting_provider_discard_candidate(diagram_input, store, change):
    deck, _, payload = diagram_input
    def mutate():
        if change == "deck":
            changed = deck.model_copy(deep=True)
            changed.meta.presenter = "동시 편집"
            store.save_deck("diagram", changed, snapshot=False)
        elif change == "deck_invalid": (store.root / "diagram" / "deck.json").write_text("invalid", encoding="utf-8")
        elif change == "source": store.write_source("diagram", "unselected.md", "변경된 원문")
        elif change == "source_too_large": store.write_source("diagram", "large.md", "자" * 100_001)
        elif change == "source_invalid": (store.root / "diagram" / "sources" / "unselected.md").write_bytes(b"\xff")
        elif change == "add_source": store.write_source("diagram", "new.md", "새 원문")
        elif change == "remove_source": (store.root / "diagram" / "sources" / "unselected.md").unlink()
        else:
            for path in (store.root / "diagram" / "sources").iterdir(): path.unlink()
    provider = StubProvider(payload, effect=mutate)
    client = app_client(store, provider)
    response = client.post(ROUTE, json=REQUEST)
    assert response.status_code == (412 if change.startswith("deck") else 409), response.text
    assert len(provider.calls) == 1
    assert store.list_snapshots("diagram") == []
    assert len(usage_records(store)) == 1
    assert client.get("/api/status").json()["last_generation_at"] is None


class FakeConnection:
    def __init__(self, provider):
        self.stub = provider
        self.status_checks = 0

    def status(self):
        self.status_checks += 1
        return LoginStatus(logged_in=True)

    def models(self): return [ModelOption(id="synthetic-model", label="합성")]
    def login_status(self): return LoginAttempt(state="idle")
    def provider(self, model): return self.stub
    def close(self): pass


def test_selection_consent_and_etag_gate_before_lease_and_release(diagram_input, store, tmp_path):
    provider = StubProvider(diagram_input[2])
    connection = FakeConnection(provider)
    manager = AIConnections(tmp_path / "ai.json", connections={"claude": connection})
    try:
        manager.select(AISelection(provider="claude", model="synthetic-model"))
        client = app_client(store, manager=manager)
        headers = {"X-AI-Selection": manager.selection_id}
        response = client.post(ROUTE, json=REQUEST, headers={**headers, "If-Match": '"old"'})
        assert response.status_code == 412
        assert connection.status_checks == 0
        response = client.post(ROUTE, json=REQUEST)
        assert response.status_code == 409
        response = client.post(ROUTE, json=REQUEST, headers={"X-AI-Selection": "old"})
        assert response.status_code == 409
        assert provider.calls == []
        def change_selection_while_busy():
            with pytest.raises(ConnectionConflict):
                manager.select(AISelection(provider="claude", model="synthetic-model"))
        provider.effect = change_selection_while_busy
        response = client.post(ROUTE, json=REQUEST, headers=headers)
        assert response.status_code == 200, response.text
        assert response.json()["status"] == "ok"
        assert len(provider.calls) == 1
        # A released generation lease allows the same selection to be explicitly applied.
        manager.select(AISelection(provider="claude", model="synthetic-model"))
    finally:
        manager.close()


@pytest.mark.parametrize("change", ["deck", "source"])
def test_changes_during_lease_acquisition_do_not_reach_provider(diagram_input, store, tmp_path, change):
    deck, _, payload = diagram_input
    provider = StubProvider(payload)
    connection = FakeConnection(provider)
    def status_after_edit():
        if change == "deck":
            changed = deck.model_copy(deep=True)
            changed.meta.presenter = "로그인 확인 중 편집"
            store.save_deck("diagram", changed, snapshot=False)
        else:
            store.write_source("diagram", "new.md", "로그인 확인 중 추가")
        return LoginStatus(logged_in=True)
    connection.status = status_after_edit
    manager = AIConnections(tmp_path / "ai.json", connections={"claude": connection})
    try:
        manager.select(AISelection(provider="claude", model="synthetic-model"))
        response = app_client(store, manager=manager).post(ROUTE, json=REQUEST, headers={
            "X-AI-Selection": manager.selection_id,
        })
        assert response.status_code == (412 if change == "deck" else 409), response.text
        assert provider.calls == []
        assert usage_records(store) == []
        manager.select(AISelection(provider="claude", model="synthetic-model"))
    finally:
        manager.close()


@pytest.mark.parametrize("defect", ["absent", "existing", "stale", "unknown_claim"])
def test_direct_service_also_rejects_invalid_context_before_provider(diagram_input, defect):
    from slidecaptain.pipeline.diagram_generation import GenerateDiagramRequest
    from slidecaptain.pipeline.service import GenerationService

    deck, sources, payload = diagram_input
    request = deepcopy(REQUEST)
    if defect == "absent": deck.structure.story_plan = None
    if defect == "existing": request["chapter_id"] = "synthetic-flow"
    if defect == "stale": deck.structure.story_plan.input_fingerprint = "0" * 64
    if defect == "unknown_claim": request["claim_ids"] = ["unknown"]
    provider = StubProvider(payload)
    service = GenerationService(provider, FontMetrics.load_default())
    records = []
    with pytest.raises(ValueError):
        asyncio.run(service.generate_diagram(deck, GenerateDiagramRequest(**request), sources, on_usage=records.append))
    assert provider.calls == []
    assert records == []


def candidate_deck(deck, diagram):
    candidate = deck.model_dump(mode="json")
    candidate["structure"]["chapters"].append({
        "id": REQUEST["chapter_id"], "topic": REQUEST["topic"], "conclusion": "", "template": "diagram", "source_refs": [],
    })
    candidate["slides"].append({
        "chapter_id": REQUEST["chapter_id"], "eyebrow": "", "subtitle": "",
        "slots": {"template": "diagram", "diagram": diagram, "footnote": "합성 후보"},
    })
    return candidate


@pytest.mark.parametrize("non_adjacent", [False, True])
def test_candidate_enters_existing_manual_validation_and_save_path(diagram_input, store, non_adjacent):
    deck, _, payload = diagram_input
    if non_adjacent:
        payload["edges"][1].update(from_node_id="intake")
    provider = StubProvider(payload)
    client = app_client(store, provider)
    result = client.post(ROUTE, json=REQUEST)
    assert result.status_code == 200, result.text
    result = result.json()
    assert result["status"] == "ok"  # Structure success does not approve a layout.
    assert result["diagram"]["edges"] == payload["edges"]
    candidate = candidate_deck(deck, result["diagram"])
    linked = client.post("/api/projects/diagram/story-plan/diagram", json={
        "deck": candidate, "chapter_id": REQUEST["chapter_id"], "role": REQUEST["role"], "claim_ids": REQUEST["claim_ids"],
    })
    assert linked.status_code == 200, linked.text
    assert store.load_deck("diagram") == deck
    render = client.post("/api/render-plan", json=linked.json())
    if non_adjacent:
        assert render.status_code == 422, render.text
        assert store.load_deck("diagram") == deck
        assert store.list_snapshots("diagram") == []
    else:
        assert render.status_code == 200, render.text
        saved = client.put("/api/projects/diagram/deck", json=linked.json())
        assert saved.status_code == 200, saved.text
        assert store.load_deck("diagram").model_dump(mode="json") == linked.json()
        assert len(store.list_snapshots("diagram")) == 1
        exported = client.post("/api/projects/diagram/export")
        assert exported.status_code == 200, exported.text
        pptx = Presentation(exported.json()["path"])
        assert len(pptx.slides) == 3
        assert any(shape.name == "new-diagram:node:intake" for shape in pptx.slides[2].shapes)
        assert exported.json()["quality"]["final_export_allowed"] is False
        assert client.post("/api/projects/diagram/export?final=true").status_code == 422
    assert len(provider.calls) == 1
