"""직접 수정의 렌더 적용과 PPTX (개정판 D3c-2, 계획 4.4, 4.6).

수정값 다섯 종의 렌더 결과, 빌더 안 글자 배율(문단과 경고가 같은 배율, 높이와 자리 불변), 한계 자르기와 읽기 하한
경고, 페이지 안 자르기, 표 묶음, 숨김과 부분 강조와 차트, PPTX 다시 읽기 대조, 수정값 없는 덱의 렌더 계획과
PPTX 바이트 불변, 출력 지문, 숨긴 근거 표시 검사, 최대 크기 라우트를 지킨다.
"""

import re
import io
import json
from pathlib import Path

import pytest
from fastapi.testclient import TestClient
from pptx import Presentation
from pptx.util import Emu, Pt

from slidecaptain.export import pptx_writer
from slidecaptain.export.pptx_writer import write_pptx
from slidecaptain.layout.direct_edits import warning_key
from slidecaptain.layout.engine import build_render_plan
from slidecaptain.metrics.font_metrics import FontMetrics
from slidecaptain.models.deck import Chapter, Deck, DeckMeta, Slide, Structure
from slidecaptain.models.direct_edit import frame_keys, numeric_path_key
from slidecaptain.models.preset import Preset
from slidecaptain.pipeline.quality import assess_quality
from slidecaptain.server.app import create_app
from slidecaptain.storage.file_store import FileProjectStore

FIXTURES = Path(__file__).parent / "fixtures"
METRICS = FontMetrics.load_default()
PRESET = Preset()


def _deck(template: str, slots: dict, *, topic="장 제목", edits=None, **slide) -> Deck:
    return Deck(meta=DeckMeta(title="보고"),
                structure=Structure(chapters=[Chapter(id="c1", topic=topic, template=template)]),
                slides=[Slide(chapter_id="c1", slots={"template": template, **slots}, direct_edits=edits or {}, **slide)])


def _plan(deck: Deck, **kw):
    return build_render_plan(deck, PRESET, METRICS, **kw)


def _parts(path) -> dict[str, bytes]:
    """PPTX(zip) 부품별 내용. 파일 바이트는 zip 항목의 쓰기 시각이 들어가 같은 내용도 다를 수 있다 (D3c-2 리뷰 F7)."""
    import zipfile

    with zipfile.ZipFile(path) as z:
        return {info.filename: z.read(info) for info in z.infolist()}


def _frames(plan) -> dict:
    return {f.name.split(":", 1)[1]: f for f in plan.slides[0].frames}


BULLETS = {"bullets": [{"text": "첫 항목"}, {"text": "둘째 항목"}], "conclusion": "결론 문장", "footnote": "출처: 사내 자료"}
CARDS = {"cards": [{"heading": "A", "badge": "1", "tail": "끝", "emphasis": True, "bullets": [{"text": "가"}]},
                   {"heading": "B", "bullets": [{"text": "나"}]}]}


# -- 위치 ---------------------------------------------------------------------------------

def test_move_adds_the_offset_records_the_engine_position_and_stays_on_the_page():
    plain = _frames(_plan(_deck("bullet_box", BULLETS)))
    moved = _frames(_plan(_deck("bullet_box", BULLETS, edits={"conclusion": {"dx_pt": 30, "dy_pt": -20},
                                                                  "title": {"dx_pt": 5000, "dy_pt": -5000}})))
    c, base = moved["conclusion"], plain["conclusion"]
    assert (c.x, c.y, c.w, c.h) == (base.x + 30, base.y - 20, base.w, base.h)
    assert (c.base_x, c.base_y) == (base.x, base.y)
    t = moved["title"]
    assert (t.x, t.y) == (PRESET.page_width_pt - t.w, 0.0)  # 페이지 안으로 자른다
    assert moved["bullets"].base_x is None  # 수정값이 없는 프레임에는 기본 자리 칸을 싣지 않는다


def test_table_group_moves_chart_and_conditions_together_and_clips_as_one():
    from test_comparison_charts_spans import chart_sample

    from slidecaptain.models.expression import ChartSpec

    deck, sources = chart_sample()
    deck.slides[0].chart = ChartSpec(rule_version="comparison-chart-v1",
                                     comparison_id=deck.structure.story_plan.comparisons[0].id, kind="column")
    plain = _frames(build_render_plan(deck, PRESET, METRICS, sources=sources))
    payload = deck.model_dump(mode="json")
    payload["slides"][0]["direct_edits"] = {"table": {"dx_pt": 0, "dy_pt": 10000}}
    deck = Deck.model_validate(payload)  # 생 dict를 대입하지 않는다(직렬화 경고, α 리뷰 A10)
    moved = _frames(build_render_plan(deck, PRESET, METRICS, sources=sources))
    dy = moved["chart"].y - plain["chart"].y
    assert dy > 0 and moved["chart_conditions"].y - plain["chart_conditions"].y == pytest.approx(dy)
    bottom = max(f.y + f.h for k, f in moved.items() if k in ("chart", "chart_conditions"))
    assert bottom == pytest.approx(PRESET.page_height_pt)


# -- 글자 크기 ------------------------------------------------------------------------------

@pytest.mark.parametrize("template,slots,key", [
    ("bullet_box", BULLETS, "bullets"), ("bullet_box", BULLETS, "conclusion"), ("bullet_box", BULLETS, "footnote"),
    ("bullet_box", BULLETS, "title"), ("cards", CARDS, "card0"), ("callout", {"text": "핵심"}, "text"),
    ("process", {"steps": [{"heading": f"단계{i}", "subtitle": "부", "notes": ["라벨"]} for i in range(3)]}, "step1"),
    ("process", {"steps": [{"heading": f"단계{i}", "notes": ["라벨"]} for i in range(3)]}, "step2_labels"),
    ("matrix", {"rows": [{"category": f"분류{i}", "primary": "대표", "items": ["가"]} for i in range(3)]}, "row0_items"),
    ("cover", {"title": "제목", "subtitle": "부제", "date": "2026"}, "date"),
    ("divider", {"section_no": "1", "section_title": "구분"}, "section_title"),
])
def test_font_scale_scales_every_paragraph_of_the_element_but_not_its_box(template, slots, key):
    """배율은 문단 크기와 줄바꿈과 경고에만 쓰고 상자 크기와 자리는 바꾸지 않는다 (재리뷰 RR2)."""
    plain = _frames(_plan(_deck(template, slots, eyebrow="" if template in ("cover", "divider") else "")))
    scaled_plan = _plan(_deck(template, slots, edits={key: {"font_scale": 1.25}}))
    scaled = _frames(scaled_plan)
    names = [n for n in plain if n == key]
    assert names
    for name in names:
        a, b = plain[name], scaled[name]
        assert (a.x, a.y, a.w, a.h) == (b.x, b.y, b.w, b.h)
        assert [p.font_pt for p in b.paras] == [round(p.font_pt * 1.25, 2) for p in a.paras]
    # 다른 요소는 그대로다
    for name, frame in plain.items():
        if name != key:
            assert [p.font_pt for p in frame.paras] == [p.font_pt for p in scaled[name].paras]


def test_paragraphs_and_warnings_use_the_same_scale():
    """문단만 키우고 경고는 기본 크기로 계산하면 넘친 상자에 경고가 없다 (재리뷰 RR26)."""
    long_title = "아주 긴 장 제목이 한 줄을 넘어 두 줄로 넘어가는지 확인한다"
    deck = _deck("bullet_box", BULLETS, topic=long_title)
    assert not [w for w in _plan(deck).slides[0].warnings if w.slot == "title"]
    big = _deck("bullet_box", BULLETS, topic=long_title, edits={"title": {"font_scale": 1.6}})
    assert [w for w in _plan(big).slides[0].warnings if w.slot == "title"]
    card = _deck("cards", CARDS, edits={"card0": {"font_scale": 1.3}})
    slots = {w.slot for w in _plan(card).slides[0].warnings}
    assert "card0_badge" in slots  # 배지 고정 높이를 넘는다(빌더의 구역 경고가 새 크기로 계산됨)


def test_font_scale_is_clamped_to_powerpoint_limits_with_a_warning():
    deck = _deck("bullet_box", BULLETS, edits={"conclusion": {"font_scale": 500}, "footnote": {"font_scale": 0.01}})
    plan = _plan(deck)
    frames = _frames(plan)
    upper = int(1584 / PRESET.spacing.line_spacing)
    assert max(p.font_pt for p in frames["conclusion"].paras) == pytest.approx(upper)
    assert min(p.font_pt for p in frames["footnote"].paras) == pytest.approx(1.0)
    messages = {w.slot: w.message for w in plan.slides[0].warnings if "맞췄습니다" in w.message}
    assert set(messages) == {"conclusion", "footnote"}


def test_shrinking_below_the_reading_floor_warns_but_is_allowed():
    """막으면 기본 크기가 하한과 같은 본문을 줄일 수 없다 (재리뷰 RR18)."""
    plan = _plan(_deck("bullet_box", BULLETS, edits={"bullets": {"font_scale": 0.9}}))
    assert min(p.font_pt for p in _frames(plan)["bullets"].paras) < 12
    [warning] = [w for w in plan.slides[0].warnings if "읽기 하한" in w.message]
    assert (warning.slot, warning.kind) == ("bullets", "font_size")


def test_font_size_warnings_are_reported_apart_from_layout_capacity():
    """일부러 줄인 글자는 분량 넘침이 아니다. 용량 항목은 통과로 두고 글자 크기 항목을 따로 보인다 (D3c-2 리뷰 F3)."""
    deck = _deck("bullet_box", BULLETS, edits={"bullets": {"font_scale": 0.9}, "footnote": {"font_scale": 0.01}})
    plan = _plan(deck)
    kinds = sorted((w.slot, w.kind) for w in plan.slides[0].warnings)
    assert kinds == [("bullets", "font_size"), ("footnote", "font_size"), ("footnote", "font_size")]  # 하한 둘과 자르기 하나
    report = assess_quality(deck, PRESET, plan)
    checks = {c.name: c for c in report.checks}
    assert checks["layout_capacity"].status == "passed" and "용량 경고 0건" in checks["layout_capacity"].detail
    assert (checks["font_size"].status, checks["font_size"].failed, checks["font_size"].chapter_ids) == ("failed", 3, ["c1"])
    plain = _deck("bullet_box", BULLETS)
    plain_plan = _plan(plain)
    assert "font_size" not in {c.name for c in assess_quality(plain, PRESET, plain_plan).checks}
    # 종류 칸은 기본값이면 직렬화에서 빠진다: 수정값이 없는 덱의 렌더 계획 바이트가 그대로다
    long = _deck("bullet_box", {**BULLETS, "conclusion": "아주 긴 결론 " * 40})
    dumped = _plan(long).model_dump(mode="json")["slides"][0]["warnings"]
    assert dumped and all("kind" not in w for w in dumped)


# -- 색 -------------------------------------------------------------------------------------

def test_text_color_paints_every_paragraph_and_fill_keeps_designed_text_color():
    plan = _plan(_deck("bullet_box", BULLETS, edits={"bullets": {"text_color": "danger"}, "conclusion": {"fill": "surface2"}}))
    frames = _frames(plan)
    assert {p.color for p in frames["bullets"].paras} == {PRESET.colors.danger}
    assert frames["conclusion"].fill == PRESET.colors.surface2
    assert {p.color for p in frames["conclusion"].paras} == {PRESET.colors.accent}  # 결론의 강조색은 그대로 (리뷰 저장 S16)


def test_fill_on_luminance_elements_recomputes_text_color_unless_text_color_is_set():
    from slidecaptain.metrics.color import readable_text_color

    c = PRESET.colors
    frames = _frames(_plan(_deck("cards", CARDS, edits={"card0": {"fill": "surface1"}})))
    expected = readable_text_color(c.surface1, light=c.background, dark=c.text)
    assert frames["card0"].fill == c.surface1 and {p.color for p in frames["card0"].paras} == {expected}
    frames = _frames(_plan(_deck("cards", CARDS, edits={"card0": {"fill": "surface1", "text_color": "accent1"}})))
    assert {p.color for p in frames["card0"].paras} == {c.accent1}


# -- 숨김 -----------------------------------------------------------------------------------

def test_hidden_frame_stays_in_the_plan_flagged_and_its_warnings_are_dropped():
    long = {**BULLETS, "conclusion": "넘치는 결론 " * 40}
    plain = _plan(_deck("bullet_box", long))
    assert any(w.slot == "conclusion" for w in plain.slides[0].warnings)
    plan = _plan(_deck("bullet_box", long, edits={"conclusion": {"hidden": True}}))
    assert _frames(plan)["conclusion"].hidden is True
    assert not any(warning_key(w.slot, frame_keys(_deck("bullet_box", long).slides[0])) == "conclusion"
                   for w in plan.slides[0].warnings)


def test_hiding_an_element_with_a_text_span_does_not_break_the_plan():
    """1판대로면 부분 강조가 숨긴 프레임을 찾지 못해 덱 전체의 렌더가 실패했다 (계획 리뷰 저장 S9)."""
    fixture = json.loads((FIXTURES / "q2d-numeric-review.json").read_text(encoding="utf-8"))
    deck = Deck.model_validate(fixture["deck"])
    slide = deck.slides[0]
    text = slide.slots.bullets[0].text
    payload = deck.model_dump(mode="json")
    payload["slides"][0]["text_spans"] = [{"slot": "bullets", "index": 0, "start": 0, "end": min(2, len(text)),
                                           "role": "bold", "text_sha256": __import__("hashlib").sha256(text.encode()).hexdigest()}]
    payload["slides"][0]["direct_edits"] = {"bullets": {"hidden": True}}
    hidden = Deck.model_validate(payload)
    plan = build_render_plan(hidden, PRESET, METRICS, sources=fixture["sources"])
    assert _frames(plan)["bullets"].hidden is True
    payload["slides"][0]["direct_edits"] = {}
    shown = build_render_plan(Deck.model_validate(payload), PRESET, METRICS, sources=fixture["sources"])
    assert _frames(shown)["bullets"].paras[0].runs  # 다시 보이면 강조가 돌아온다


def test_hidden_table_skips_chart_construction_even_when_the_chart_cannot_be_built():
    from test_comparison_charts_spans import chart_sample

    from slidecaptain.models.expression import ChartSpec

    deck, sources = chart_sample()
    deck.slides[0].chart = ChartSpec(rule_version="comparison-chart-v1",
                                     comparison_id=deck.structure.story_plan.comparisons[0].id, kind="column")
    with pytest.raises(Exception):
        build_render_plan(deck, PRESET, METRICS, sources={})  # 자료가 없으면 차트를 만들 수 없다
    payload = deck.model_dump(mode="json")
    payload["slides"][0]["direct_edits"] = {"table": {"hidden": True}}
    deck = Deck.model_validate(payload)
    plan = build_render_plan(deck, PRESET, METRICS, sources={})
    frames = _frames(plan)
    assert frames["table"].hidden is True and "chart" not in frames


# -- 직렬화와 PPTX ---------------------------------------------------------------------------

def _sample_decks():
    for path in sorted(FIXTURES.glob("*.json")):
        payload = json.loads(path.read_text(encoding="utf-8"))
        deck = payload.get("deck") if isinstance(payload, dict) else None
        if isinstance(deck, dict) and deck.get("slides"):
            try:
                yield path.name, Deck.model_validate(deck), payload.get("sources") or {}
            except Exception:
                continue


def test_plans_without_edits_have_no_new_fields():
    for name, deck, sources in _sample_decks():
        try:
            plan = build_render_plan(deck, PRESET, METRICS, sources=sources)
        except Exception:
            continue
        text = plan.model_dump_json()
        assert '"base_x"' not in text and '"hidden"' not in text, name


def test_pptx_bytes_are_unchanged_for_decks_without_edits(tmp_path, monkeypatch):
    """글자 크기는 1/100pt 반올림으로 바꿨지만 기본 프리셋의 크기에서는 버림과 같아 바이트가 그대로다 (재리뷰 RR1).

    덱별 디자인 값에 18.4pt 같은 소수 크기를 넣은 덱은 바뀐다. 아래 시험이 그 범위를 고정한다 (D3c-2 리뷰 F1)."""
    original = pptx_writer._style_run

    def floor_style_run(run, para, style):
        original(run, para, style)
        run.font.size = Pt(para.font_pt)

    checked = 0
    for name, deck, sources in _sample_decks():
        try:
            plan = build_render_plan(deck, PRESET, METRICS, sources=sources)
        except Exception:
            continue
        new = tmp_path / f"{name}.new.pptx"
        write_pptx(plan, new)
        monkeypatch.setattr(pptx_writer, "_style_run", floor_style_run)
        old = tmp_path / f"{name}.old.pptx"
        write_pptx(plan, old)
        monkeypatch.setattr(pptx_writer, "_style_run", original)
        assert _parts(new) == _parts(old), name  # zip 항목 시각(2초 단위)은 비교하지 않는다 (D3c-2 리뷰 F7)
        checked += 1
    assert checked >= 2


def test_decimal_design_sizes_are_written_rounded_not_floored(tmp_path):
    """수정값이 없어도 디자인 값의 소수 크기는 반올림으로 기록돼 0.2.0의 버림보다 0.01pt 크다 (D3c-2 리뷰 F1).

    그 덱의 PPTX 바이트는 이 개정판에서 바뀐다. 내보내기 지문은 덱과 렌더 계획으로 계산하므로 바뀌지 않는다."""
    preset = PRESET.model_copy(update={"font_roles": PRESET.font_roles.model_copy(update={"title_pt": 18.4, "footnote_pt": 9.2})})
    deck = _deck("bullet_box", BULLETS)
    out = tmp_path / "decimal.pptx"
    write_pptx(build_render_plan(deck, preset, METRICS), out)
    xml = Presentation(str(out)).slides[0].shapes._spTree.xml
    assert 'sz="1840"' in xml and 'sz="920"' in xml
    assert 'sz="1839"' not in xml and 'sz="919"' not in xml


def test_pptx_matches_the_plan_for_edited_elements(tmp_path):
    deck = _deck("bullet_box", BULLETS, edits={
        # 12pt x 0.8575 = 10.29pt. Pt의 버림은 1028(10.28pt)로 기록한다 (계획 리뷰 저장 S8의 실측 값)
        "conclusion": {"dx_pt": 33.3, "dy_pt": -12.7, "font_scale": 0.8575, "text_color": "danger", "fill": "surface3"},
        "footnote": {"hidden": True},
    })
    plan = _plan(deck)
    path = tmp_path / "edited.pptx"
    write_pptx(plan, path)
    slide = Presentation(path).slides[0]
    shapes = {s.name: s for s in slide.shapes}
    assert "c1:footnote" not in shapes  # 숨긴 도형은 쓰지 않는다
    frame = _frames(plan)["conclusion"]
    shape = shapes["c1:conclusion"]
    assert (shape.left, shape.top) == (Emu(round(frame.x * 12700)), Emu(round(frame.y * 12700)))
    runs = [r for p in shape.text_frame.paragraphs for r in p.runs]
    assert frame.paras[0].font_pt == 10.29
    assert {r.font.size.centipoints for r in runs} == {1029}
    assert {str(r.font.color.rgb) for r in runs} == {PRESET.colors.danger.upper()}
    assert str(shape.fill.fore_color.rgb) == PRESET.colors.surface3.upper()


# -- 지문과 검토 ------------------------------------------------------------------------------

def test_moving_an_element_changes_the_export_fingerprint_only_when_edited():
    deck = _deck("bullet_box", BULLETS)
    plain = assess_quality(deck, PRESET, _plan(deck))
    again = assess_quality(deck, PRESET, _plan(deck))
    moved_deck = _deck("bullet_box", BULLETS, edits={"title": {"dx_pt": 1, "dy_pt": 0}})
    moved = assess_quality(moved_deck, PRESET, _plan(moved_deck))
    assert plain.input_fingerprint == again.input_fingerprint != moved.input_fingerprint


def test_hidden_evidence_check_appears_only_when_evidence_is_hidden():
    deck = _deck("bullet_box", BULLETS, eyebrow="분류", edits={"eyebrow": {"hidden": True}})
    names = [c.name for c in assess_quality(deck, PRESET, _plan(deck)).checks]
    assert "hidden_evidence" not in names  # 근거가 아닌 요소의 숨김은 검사를 넣지 않는다
    deck = _deck("bullet_box", BULLETS, edits={"footnote": {"hidden": True}})
    report = assess_quality(deck, PRESET, _plan(deck))
    [check] = [c for c in report.checks if c.name == "hidden_evidence"]
    assert (check.status, check.failed, check.chapter_ids) == ("failed", 1, ["c1"])
    assert report.status == "needs_revision"


def test_hidden_numeric_text_is_evidence_even_when_numeric_review_did_not_run():
    # 구성 계획이 없는 덱: 수치 검토는 not_run(missing_plan)이고 평가 항목이 비지만 숫자 문구의 숨김은 잡는다 (D3c-2 리뷰 F2)
    deck = _deck("bullet_box", {"bullets": [{"text": "매출 20% 증가"}, {"text": "고객 만족"}], "conclusion": "영업이익 35억 원 달성"},
                 edits={"bullets": {"hidden": True}, "conclusion": {"hidden": True}})
    report = assess_quality(deck, PRESET, _plan(deck))
    assert (report.numeric_review.status, report.numeric_review.reason) == ("not_run", "missing_plan")
    [check] = [c for c in report.checks if c.name == "hidden_evidence"]
    assert check.failed == 2
    # 장 번호, 장 제목, 우리말 요소 이름으로 적고 내부 열쇠는 쓰지 않는다 (D3c-2 리뷰 F6)
    assert "1장 장 제목의 본문 불릿, 결론 상자" in check.detail and "bullets" not in check.detail
    plain = _deck("bullet_box", {"bullets": [{"text": "고객 만족"}, {"text": "품질 개선"}], "conclusion": "결론"}, edits={"bullets": {"hidden": True}})
    assert "hidden_evidence" not in [c.name for c in assess_quality(plain, PRESET, _plan(plain)).checks]


def test_numeric_paths_map_to_element_keys():
    cover = Slide(chapter_id="c1", slots={"template": "cover", "title": "t", "subtitle": "s"}, eyebrow="e")
    assert numeric_path_key(cover, "/slots/title") == "cover_title"
    assert numeric_path_key(cover, "/slots/subtitle") == "subtitle"
    assert numeric_path_key(cover, "/subtitle") is None and numeric_path_key(cover, "/eyebrow") is None
    table = Slide(chapter_id="c1", slots={"template": "table", "columns": ["a", "b"], "rows": [["1", "2"]], "footnote": "f"})
    assert numeric_path_key(table, "/slots/rows/0/1") == "table" and numeric_path_key(table, "/slots/footnote") == "footnote"
    process = _deck("process", {"steps": [{"heading": "h", "subtitle": "s", "notes": ["n"]} for _ in range(3)]}).slides[0]
    assert numeric_path_key(process, "/slots/steps/1/notes/0") == "step1_labels"
    assert numeric_path_key(process, "/slots/steps/1/subtitle") == "step1"
    matrix = _deck("matrix", {"rows": [{"category": "c", "primary": "p", "items": ["i"]} for _ in range(3)]}).slides[0]
    assert numeric_path_key(matrix, "/slots/rows/2/items/0") == "row2_items"
    assert numeric_path_key(matrix, "/slots/rows/2/primary") == "row2_primary"
    assert numeric_path_key(_deck("bullet_box", BULLETS).slides[0], "/topic") == "title"


def test_warning_slots_map_to_element_keys():
    keys = frame_keys(_deck("process", {"steps": [{"heading": "h", "notes": ["n"]} for _ in range(3)]}).slides[0])
    assert warning_key("step0_labels_1", keys) == "step0_labels"
    assert warning_key("step0_heading", keys) == "step0"
    assert warning_key("step0_badge", keys) == "step0_badge"
    cards = frame_keys(_deck("cards", CARDS).slides[0])
    assert warning_key("card0_badge", cards) == "card0" and warning_key("card0_tail", cards) == "card0"
    assert warning_key("bullets_span", frame_keys(_deck("bullet_box", BULLETS).slides[0])) == "bullets"
    assert warning_key("chart", {"table"}) == "table"


# -- 최대 크기 라우트 ---------------------------------------------------------------------------

@pytest.fixture
def client(tmp_path):
    store = FileProjectStore(tmp_path / "projects")
    store.create_project("p1")
    with TestClient(create_app(store), headers={"X-Requested-With": "SlideCaptain"}) as c:
        yield c


def test_font_limit_reports_the_largest_size_that_does_not_overflow(client):
    deck = _deck("bullet_box", BULLETS, eyebrow="분류")
    body = {"deck": deck.model_dump(mode="json"), "chapter_id": "c1", "key": "eyebrow"}
    limit = client.post("/api/projects/p1/direct-edits/font-limit", json=body).json()
    assert limit["current_pt"] == limit["base_pt"] == PRESET.font_roles.eyebrow_pt
    # 아이브로우는 높이가 한 줄에 딱 맞아 키울 여유가 없다 (재리뷰 RR2 실측: 최대 배율 1.000)
    assert limit["fits_up_to_pt"] == pytest.approx(PRESET.font_roles.eyebrow_pt, abs=0.05)
    assert limit["min_pt"] == 1.0 and limit["max_pt"] > 100
    body["key"] = "title"
    title = client.post("/api/projects/p1/direct-edits/font-limit", json=body).json()
    assert title["fits_up_to_pt"] > PRESET.font_roles.title_pt


def test_font_limit_counts_bold_span_width_like_the_engine():
    """굵게 강조한 요소도 보인 최대 크기로 저장하면 그 요소에 경고가 없다 (D3c-2 리뷰 F4, 리뷰 실측 8개 중 4개 경고)."""
    import hashlib

    from slidecaptain.layout.direct_edits import font_limit

    warned = []
    for n in range(1, 9):
        first = ("가나다라마바사 아자차카 타파하 " * n).strip()
        span = {"slot": "bullets", "index": 0, "start": 0, "end": len(first), "role": "bold",
                "text_sha256": hashlib.sha256(first.encode()).hexdigest()}
        deck = _deck("bullet_box", {"bullets": [{"text": first}, {"text": "둘째 항목 문장"}], "conclusion": "결론"},
                     text_spans=[span])
        limit = font_limit(deck.structure.chapters[0], deck.slides[0], PRESET, METRICS, "", "bullets")
        if limit["fits_up_to_pt"] is None:
            continue
        payload = deck.model_dump(mode="json")
        payload["slides"][0]["direct_edits"] = {"bullets": {"font_scale": limit["fits_up_to_pt"] / limit["base_pt"]}}
        edited = Deck.model_validate(payload)
        keys = frame_keys(edited.slides[0])
        if any(warning_key(w.slot, keys) == "bullets" for w in _plan(edited).slides[0].warnings):
            warned.append(n)
    assert warned == []


def test_font_limit_rejects_elements_without_font_and_unknown_chapters(client):
    deck = _deck("table", {"columns": ["a"], "rows": [["1"]]})
    body = {"deck": deck.model_dump(mode="json"), "chapter_id": "c1", "key": "table"}
    assert client.post("/api/projects/p1/direct-edits/font-limit", json=body).status_code == 422
    body["chapter_id"] = "없음"
    assert client.post("/api/projects/p1/direct-edits/font-limit", json=body).status_code == 404


def test_warning_keys_fixture_matches_the_engine():
    """화면의 warningKey(D3c-3)가 같은 결과를 내는지 지키는 고정 파일. 모든 템플릿의 실제 경고 slot을 담는다."""
    cases = json.loads((FIXTURES / "warning-keys.json").read_text(encoding="utf-8"))["cases"]
    assert len(cases) > 100
    for case in cases:
        assert warning_key(case["slot"], set(case["keys"])) == case["key"], case
    assert [c["slot"] for c in cases if c["key"] is None] == ["unknown_slot"]


# -- D3c-2 리뷰 F8~F11: 변형이 살아남은 자리의 시험 ----------------------------------------------------

def _lengthen(value):
    """분량 경고가 나도록 문구를 늘인다(글자 크기와 상관없는 칸은 그대로)."""
    if isinstance(value, str):
        return (value + " 길게 늘여 쓴 문구입니다") * 6 if value.strip() else value
    if isinstance(value, list):
        return [_lengthen(v) for v in value]
    if isinstance(value, dict):
        return {k: v if k in ("template", "tone") else _lengthen(v) for k, v in value.items()}
    return value


def _sample_slides(long: bool):
    from pydantic import ValidationError
    from test_direct_edits import SAMPLES

    for template in SAMPLES:
        for variant in ("full", "bare"):
            slots = {"template": template, **SAMPLES[template][variant]}
            try:
                yield Slide(chapter_id="c1", slots=_lengthen(slots) if long else slots)
            except ValidationError:
                continue


def _element(page, key):
    """한 요소의 문단(글, 크기, 줄)과 그 요소에 대응하는 경고."""
    from slidecaptain.layout.direct_edits import frame_key

    keys = {frame_key(f) for f in page.frames}
    frames = [(f.name, [(p.text, p.font_pt, p.lines) for p in f.paras]) for f in page.frames if frame_key(f) == key]
    warnings = sorted((w.slot, round(w.needed_pt, 3), round(w.available_pt, 3))
                      for w in page.warnings if warning_key(w.slot, keys) == key)
    return frames, warnings


@pytest.mark.parametrize("long", [False, True])
def test_element_scale_matches_a_globally_scaled_preset(long):
    """요소별 배율의 문단과 경고는 프리셋 전체에 같은 배율을 준 결과와 같다 (계획 4.4, D3c-2 리뷰 F8의 M1~M4).

    공통 슬롯은 기하가 글자 크기에 기대므로 뺀다(최대 크기 라우트 시험이 덮는다)."""
    from slidecaptain.layout.templates import _font_preset, build_slide
    from slidecaptain.metrics.capacity import process_geometry

    checked = 0
    for slide in _sample_slides(long):
        chapter = Chapter(id="c1", topic="장 제목", template=slide.slots.template)
        for key, spec in frame_keys(slide).items():
            if "font" not in spec.kinds:
                continue
            for scale in (0.8, 1.3):
                own = build_slide(chapter, slide.slots, 1, PRESET, METRICS, presenter="", font_scales={key: scale})
                whole = build_slide(chapter, slide.slots, 1, _font_preset(PRESET, scale), METRICS, presenter="")
                expected = _element(whole, key)
                if slide.slots.template == "process" and re.fullmatch(r"step\d+", key):
                    # 단계의 부제 자리는 보조 라벨 줄 높이(본문 글자 크기)를 뺀 나머지다(2026-10-18 C3). 프리셋 전체 배율은
                    # 라벨도 줄여 부제 자리를 넓히지만, 단계 글자만 키우거나 줄이면 라벨은 그대로라 부제 자리도 그대로다
                    subtitle_h = round(process_geometry(PRESET, len(slide.slots.steps))["subtitle_h"], 3)
                    expected = (expected[0], [(slot, needed, subtitle_h if slot.endswith("_subtitle") else available)
                                              for slot, needed, available in expected[1]])
                if slide.slots.template == "process" and re.fullmatch(r"step\d+_labels", key):
                    # 라벨 칸의 높이는 같은 행의 단계 글자가 끝나는 자리에 기댄다. 그래서 라벨만 키운 결과와 전체를 키운
                    # 결과의 경고는 다르다(아래 test_scaled_process_note_warns_when_it_leaves_the_row). 문단만 맞춘다
                    assert _element(own, key)[0] == expected[0], (slide.slots.template, key, scale)
                else:
                    assert _element(own, key) == expected, (slide.slots.template, key, scale)
                checked += 1
    assert checked > 40


def test_scaled_process_note_warns_when_it_leaves_the_row():
    """6단계 행에 부제와 라벨이 있으면 라벨 칸에 기본 크기 한 줄(16.8pt)만 남는다. 라벨만 1.3배로 키우면 그 줄(21.84pt)이
    행 밖으로 나가므로 넘침 경고를 낸다. 1배는 경고가 없다 (2026-10-18 C3, 라벨은 배율 받은 줄 높이로 검사한다)."""
    from test_direct_edits import SAMPLES
    from slidecaptain.layout.templates import build_slide

    slide = Slide(chapter_id="c1", slots={"template": "process", **SAMPLES["process"]["full"]})
    chapter = Chapter(id="c1", topic="장 제목", template="process")
    for scale, expected in ((1.0, []), (1.3, [("step0_labels_0", 21.84, 16.8)])):
        page = build_slide(chapter, slide.slots, 1, PRESET, METRICS, presenter="", font_scales={"step0_labels": scale})
        labels = [(w.slot, round(w.needed_pt, 2), round(w.available_pt, 2)) for w in page.warnings if "_labels" in w.slot]
        assert labels == expected, scale


def test_hiding_an_element_drops_every_warning_of_that_element():
    """카드를 숨기면 card0_heading, 단계 본문을 숨기면 step0_heading과 step0_subtitle 경고도 사라진다 (F8의 M7, M12)."""
    warned = 0
    for slide in _sample_slides(True):
        deck = Deck(meta=DeckMeta(title="보고"),
                    structure=Structure(chapters=[Chapter(id="c1", topic="장 제목", template=slide.slots.template)]), slides=[slide])
        plan = _plan(deck)
        keys = frame_keys(slide)
        for key, spec in keys.items():
            if "hide" not in spec.kinds or not any(warning_key(w.slot, keys) == key for w in plan.slides[0].warnings):
                continue
            warned += 1
            payload = deck.model_dump(mode="json")
            payload["slides"][0]["direct_edits"] = {key: {"hidden": True}}
            hidden = _plan(Deck.model_validate(payload))
            assert not [w.slot for w in hidden.slides[0].warnings if warning_key(w.slot, keys) == key], key
    assert warned >= 10


def test_fill_recomputes_text_color_for_cards_badges_and_categories():
    """채움을 바꾸면 카드, 단계 번호, 분류 칸의 글자색을 채움에 맞춰 다시 정한다. 두 방향을 본다 (F8의 M6)."""
    from test_direct_edits import SAMPLES

    from slidecaptain.metrics.color import readable_text_color

    for template, key in (("cards", "card0"), ("process", "step0_badge"), ("matrix", "row0_category")):
        colors = set()
        for fill in ("ink", "surface1"):
            plan = _plan(_deck(template, SAMPLES[template]["full"], edits={key: {"fill": fill}}))
            expected = readable_text_color(getattr(PRESET.colors, fill), light=PRESET.colors.background, dark=PRESET.colors.text)
            assert {p.color for p in _frames(plan)[key].paras} == {expected}, (template, fill)
            colors.add(expected)
        assert len(colors) == 2


def test_footnote_below_its_own_floor_warns():
    """각주는 각주 하한(9pt)으로 본다 (F8의 M11)."""
    plan = _plan(_deck("bullet_box", BULLETS, edits={"footnote": {"font_scale": 0.9}}))
    [warning] = [w for w in plan.slides[0].warnings if w.slot == "footnote"]
    assert "읽기 하한(9pt)" in warning.message and warning.kind == "font_size"


def test_mixed_size_elements_clamp_by_their_smallest_and_largest_paragraph(tmp_path):
    """배지(작은 글자)와 소제목이 섞인 카드에서 하한은 가장 작은 문단, 상한은 가장 큰 문단으로 자른다. 한계의 PPTX를 쓰고
    다시 연다 (F8의 M13, M14, F9의 "한계 근처 글자 크기의 내보내기 성공", F11의 문구)."""
    for scale, check in ((1e-6, "low"), (1e6, "high")):
        deck = _deck("cards", CARDS, edits={"card0": {"font_scale": scale}})
        plan = _plan(deck)
        sizes = [p.font_pt for p in _frames(plan)["card0"].paras]
        assert len(set(sizes)) > 1  # 크기가 섞인 요소다
        if check == "low":
            assert min(sizes) == pytest.approx(1.0)
            [warning] = [w for w in plan.slides[0].warnings if w.slot == "card0" and "맞췄습니다" in w.message]
            assert warning.message.startswith("가장 작은 글자를 1pt로 맞췄습니다")
        else:
            assert max(sizes) * PRESET.spacing.line_spacing <= 1584
        out = tmp_path / f"{check}.pptx"
        write_pptx(plan, out)
        assert Presentation(str(out)).slides[0].shapes


def test_text_color_applies_before_text_spans():
    """글자 색을 바꾼 요소의 강조 밖 글은 새 색, 강조색 글은 강조색이다 (F8의 M15)."""
    import hashlib

    first = "앞부분 강조 뒷부분"
    span = {"slot": "bullets", "index": 0, "start": 4, "end": 6, "role": "accent",
            "text_sha256": hashlib.sha256(first.encode()).hexdigest()}
    deck = _deck("bullet_box", {"bullets": [{"text": first}], "conclusion": "결론"}, text_spans=[span],
                 edits={"bullets": {"text_color": "danger"}})
    runs = _frames(_plan(deck))["bullets"].paras[0].runs
    assert {(r.text, r.color) for r in runs} == {
        ("앞부분 ", PRESET.colors.danger), ("강조", PRESET.colors.accent1), (" 뒷부분", PRESET.colors.danger)}


def test_empty_cover_fields_take_no_font_scale(client):
    """빈 표지 부제는 글이 없어 배율을 버리고, 최대 크기 라우트는 422다 (D3c-2 리뷰 F10)."""
    deck = _deck("cover", {"title": "제목", "subtitle": ""}, edits={"subtitle": {"font_scale": 3}})
    plan = _plan(deck)
    assert {p.font_pt for p in _frames(plan)["subtitle"].paras} == {PRESET.font_roles.subtitle_pt}
    body = {"deck": deck.model_dump(mode="json"), "chapter_id": "c1", "key": "subtitle"}
    assert client.post("/api/projects/p1/direct-edits/font-limit", json=body).status_code == 422
    body["key"] = "cover_title"
    assert client.post("/api/projects/p1/direct-edits/font-limit", json=body).status_code == 200


# -- α 리뷰 반영 ------------------------------------------------------------------------------

def test_card_order_badges_are_not_evidence_and_other_hidden_elements_are_listed():
    """순번 배지만 숫자인 카드의 숨김은 근거 실패가 아니다(A5). 근거가 아닌 숨김은 미수행 정보 줄로 보인다(A6)."""
    cards = {"cards": [{"heading": "가", "badge": "1"}, {"heading": "나", "badge": "2"}]}
    deck = _deck("cards", cards, edits={"card0": {"hidden": True}})
    report = assess_quality(deck, PRESET, _plan(deck))
    checks = {c.name: c for c in report.checks}
    assert "hidden_evidence" not in checks and report.status == "draft"
    assert checks["hidden_elements"].status == "not_run"
    assert checks["hidden_elements"].detail.startswith("숨긴 요소 1개: 1장 장 제목의 1번 카드.")
    numeric = _deck("cards", {"cards": [{"heading": "매출 20% 증가", "badge": "1"}, {"heading": "나"}]},
                    edits={"card0": {"hidden": True}})
    assert "hidden_evidence" in {c.name for c in assess_quality(numeric, PRESET, _plan(numeric)).checks}
    plain = _deck("cards", cards)
    assert "hidden_elements" not in {c.name for c in assess_quality(plain, PRESET, _plan(plain)).checks}


def test_hidden_numeric_compare_card_is_evidence():
    """비교 카드의 숫자 문구 경로가 left_card로 간다 (α 리뷰 A8의 B04)."""
    slots = {"left": {"heading": "직접 진출", "bullets": [{"text": "비용 30% 절감"}]}, "right": {"heading": "제휴"},
             "conclusion": "결론"}
    deck = _deck("compare2", slots, edits={"left_card": {"hidden": True}})
    [check] = [c for c in assess_quality(deck, PRESET, _plan(deck)).checks if c.name == "hidden_evidence"]
    assert "왼쪽 카드" in check.detail


def test_hidden_element_drops_its_span_warning():
    """숨긴 요소의 부분 강조 경고도 사라진다. 숨김의 경고 제거는 강조 적용 뒤다 (α 리뷰 A8의 B09)."""
    import hashlib

    first = "가나다라마바사 아자차카 타파하 " * 10
    span = {"slot": "bullets", "index": 0, "start": 0, "end": len(first), "role": "bold",
            "text_sha256": hashlib.sha256(first.encode()).hexdigest()}
    slots = {"bullets": [{"text": first}, {"text": "둘째"}], "conclusion": "결론"}
    shown = _plan(_deck("bullet_box", slots, text_spans=[span], edits={"bullets": {"font_scale": 3}}))
    assert "bullets_span" in {w.slot for w in shown.slides[0].warnings}  # 전제: 숨기지 않으면 강조 경고가 있다
    hidden = _plan(_deck("bullet_box", slots, text_spans=[span], edits={"bullets": {"font_scale": 3, "hidden": True}}))
    assert not [w for w in hidden.slides[0].warnings if w.slot.startswith("bullets")]


def test_position_is_clipped_at_the_left_and_top_page_edges():
    """가로와 세로 모두 왼쪽, 위 경계에서 자른다 (α 리뷰 A8의 B11)."""
    frames = _frames(_plan(_deck("bullet_box", BULLETS, edits={"title": {"dx_pt": -5000, "dy_pt": -5000}})))
    assert (frames["title"].x, frames["title"].y) == (0, 0)


def test_font_limit_route_reports_current_base_and_deck_design_sizes(client):
    """최대 크기 라우트의 계약: 지금 배율의 현재 크기, 가장 큰 문단의 기준 크기, 덱별 디자인 값 (α 리뷰 A8의 B14, B15, B23)."""
    def ask(deck, key):
        body = {"deck": deck.model_dump(mode="json"), "chapter_id": "c1", "key": key}
        return client.post("/api/projects/p1/direct-edits/font-limit", json=body).json()

    scaled = ask(_deck("bullet_box", BULLETS, edits={"bullets": {"font_scale": 1.25}}), "bullets")
    assert (scaled["base_pt"], scaled["current_pt"]) == (PRESET.font_roles.body_pt, round(PRESET.font_roles.body_pt * 1.25, 2))
    mixed = ask(_deck("cards", CARDS), "card0")  # 배지(작은 글자)와 소제목이 섞인 카드
    assert mixed["base_pt"] == max(p.font_pt for p in _frames(_plan(_deck("cards", CARDS)))["card0"].paras)
    designed = _deck("bullet_box", BULLETS)
    designed.meta.preset_overrides = {"font_roles": {"body_pt": 14}}
    assert ask(Deck.model_validate(designed.model_dump(mode="json")), "bullets")["base_pt"] == 14
