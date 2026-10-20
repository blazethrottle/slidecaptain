"""Missing target PowerPoint can never produce a successful render receipt."""

import pytest
import threading
import time
from concurrent.futures import ThreadPoolExecutor


def test_missing_native_runtime_fails_closed(tmp_path, monkeypatch):
    from slidecaptain.export import renderer
    monkeypatch.setattr(renderer.sys, 'platform', 'darwin')
    with pytest.raises(renderer.RenderUnavailable, match='PowerPoint'):
        renderer.render_powerpoint(tmp_path/'input.pptx', tmp_path/'pages', 1)


def test_png_validation_rejects_unusable_image():
    from slidecaptain.export.renderer import validate_png
    with pytest.raises(ValueError):
        validate_png(b'not png')


def test_different_project_renders_cannot_enter_native_singleton_together(tmp_path, monkeypatch):
    from slidecaptain.export import renderer
    active = maximum = 0
    guard = threading.Lock()
    monkeypatch.setattr(renderer, '_runtime_paths', lambda: (tmp_path/'PowerPoint', tmp_path/'powershell'))
    def simulated_native(*args):
        nonlocal active, maximum
        with guard:
            active += 1
            maximum = max(maximum, active)
        time.sleep(0.03)
        with guard:
            active -= 1
        return {}, []
    monkeypatch.setattr(renderer, '_render_powerpoint_unlocked', simulated_native)
    with ThreadPoolExecutor(max_workers=2) as pool:
        results = [pool.submit(renderer.render_powerpoint,tmp_path/f'{n}.pptx',tmp_path/f'{n}',1) for n in range(2)]
        assert all(result.result() == ({},[]) for result in results)
    assert maximum == 1


def test_native_adapter_never_quits_an_unowned_user_application():
    from slidecaptain.export import renderer
    assert '.Quit(' not in renderer._SCRIPT
    assert '.AutomationSecurity' not in renderer._SCRIPT


def test_native_adapter_waits_for_its_own_powerpoint_without_killing_it():
    # 회사 PC 이월표 (9): 렌더가 COM 참조만 놓고 돌아와, 남아 있던 PowerPoint 때문에 바로 이어진 렌더가 "실행 중인
    # PowerPoint"로 거부됐다. 렌더가 띄운 프로세스가 끝나기를 기다리되, 강제로 끝내지 않는다
    from slidecaptain.export import renderer
    assert 'Wait-Process' in renderer._SCRIPT
    assert 'Stop-Process' not in renderer._SCRIPT and 'Kill(' not in renderer._SCRIPT
    # 남은 COM 참조를 거둬야 PowerPoint가 곧 끝난다. 이 줄이 빠지면 기다림이 매번 제한까지 간다(리뷰 DPR-08, 추정)
    assert '[GC]::Collect(); [GC]::WaitForPendingFinalizers()' in renderer._SCRIPT


def test_python_timeout_covers_the_render_and_the_exit_wait():
    # 리뷰 DPR-01: 스크립트는 렌더 뒤 PowerPoint 종료를 최대 30초 기다린다. 파이썬 제한이 렌더 120초만 잡으면, 오래
    # 걸린 렌더가 페이지를 다 만들고도 시간 초과로 기록될 수 있다
    import re
    from slidecaptain.export import renderer
    wait = int(re.search(r'Wait-Process -Id \$owned -Timeout (\d+)', renderer._SCRIPT).group(1))
    assert renderer.RENDER_TIMEOUT_SECONDS >= 120 + wait


def _powerpoint_ready():
    import os
    import subprocess
    import sys
    if sys.platform != 'win32' or os.environ.get('SLIDECAPTAIN_REAL_POWERPOINT') != '1':
        return False
    running = subprocess.run(['tasklist', '/FI', 'IMAGENAME eq POWERPNT.EXE', '/NH'], capture_output=True, text=True)
    return 'POWERPNT' not in running.stdout.upper()


@pytest.mark.skipif(not _powerpoint_ready(), reason='실제 PowerPoint가 있고 실행 중이 아닌 Windows에서 '
                    'SLIDECAPTAIN_REAL_POWERPOINT=1로만 실행한다')
def test_two_consecutive_real_renders_both_succeed(tmp_path):
    # 회사 PC 이월표 (9)의 실측 재현. 사용자 PowerPoint가 열려 있으면 건너뛴다(그 창을 건드리지 않는다)
    from pptx import Presentation
    from slidecaptain.export import renderer
    deck = Presentation()
    deck.slides.add_slide(deck.slide_layouts[6])
    deck.save(tmp_path / 'one.pptx')
    for attempt in range(2):
        _environment, pages = renderer.render_powerpoint(tmp_path / 'one.pptx', tmp_path / f'pages{attempt}', 1)
        assert len(pages) == 1
