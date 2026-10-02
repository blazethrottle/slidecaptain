"""동일 계산의 화면/PPTX 소비 계약. 목표 PowerPoint 표시 검사는 아니다."""

from copy import deepcopy
import json
from pathlib import Path

import pytest
from pptx import Presentation
from pptx.enum.text import MSO_AUTO_SIZE
from pptx.oxml.ns import qn

from slidecaptain.export.pptx_writer import write_pptx
from slidecaptain.layout.diagram_page import build_diagram_render_plan, DiagramRenderBlocked
from slidecaptain.layout.templates import slide_geometry
from slidecaptain.metrics.font_metrics import FontMetrics
from slidecaptain.models.preset import Preset
from slidecaptain.models.render import RenderPlan
from slidecaptain.models.story import Evidence


@pytest.fixture
def sample():
    return json.loads((Path(__file__).parent / 'fixtures/q3a-diagram.json').read_text('utf-8'))


def build(sample, preset=None, **kwargs):
    return build_diagram_render_plan(
        sample['diagram'], preset or Preset(), FontMetrics.load_default(),
        evidence=[Evidence.model_validate(e) for e in sample['evidence']],
        **({'title': '요청 처리 흐름', 'eyebrow': '합성 운영', 'subtitle': '접수부터 검토까지',
            'footnote': '합성 자료를 사용한 렌더 검사'} | kwargs),
    )


def test_common_slots_reserve_space_and_keep_all_meaning(sample):
    before = deepcopy(sample)
    plan = build(sample)
    page = plan.slides[0].diagram
    assert page is not None
    assert not plan.slides[0].frames
    assert page.layout.content_bounds.y == slide_geometry(Preset(), '합성 운영', '접수부터 검토까지')['content_top']
    headers = {f.name.rsplit(':', 1)[-1]: f for f in page.headers}
    assert set(headers) == {'title', 'eyebrow', 'subtitle', 'footnote', 'page_number'}
    assert headers['subtitle'].y + headers['subtitle'].h < page.layout.content_bounds.y
    assert page.layout.content_bounds.y + page.layout.content_bounds.h <= headers['footnote'].y
    assert [n.node.model_dump() for n in page.layout.nodes] == sample['diagram']['nodes']
    assert [e.edge.model_dump() for e in page.layout.edges] == sample['diagram']['edges']
    assert page.review.model_dump() == dict.fromkeys(['semantic', 'browser', 'powerpoint', 'reader'], 'not_run')
    assert build(sample) == plan
    assert RenderPlan.model_validate_json(plan.model_dump_json()) == plan
    assert before == sample


@pytest.mark.parametrize('field', ['title', 'eyebrow', 'subtitle', 'footnote'])
def test_every_common_text_is_fingerprinted(sample, field):
    old = build(sample).slides[0].diagram.input_fingerprint
    assert build(sample, **{field: '변경 문구'}).slides[0].diagram.input_fingerprint != old


def test_page_number_is_fingerprinted_and_must_fit(sample):
    assert build(sample, page_no=2).slides[0].diagram.input_fingerprint != build(sample).slides[0].diagram.input_fingerprint
    for number in (0, -1, True, '1'):
        with pytest.raises(ValueError):
            build(sample, page_no=number)
    with pytest.raises(DiagramRenderBlocked):
        build(sample, page_no=10**30)


def test_common_geometry_and_unrepresentable_spacing_are_blocked(sample):
    preset = Preset()
    preset.spacing.page_number_bottom = 70
    with pytest.raises(DiagramRenderBlocked):
        build(sample, preset)
    preset.spacing.page_number_bottom = 5
    with pytest.raises(DiagramRenderBlocked):
        build(sample, preset)
    preset = Preset()
    preset.spacing.line_spacing = 1.333
    with pytest.raises(DiagramRenderBlocked):
        build(sample, preset)


@pytest.mark.parametrize('font', [0.1, 0.5, 4000.01])
def test_pptx_font_range_is_checked_before_computing(sample, font):
    preset = Preset()
    preset.font_roles.title_pt = font
    with pytest.raises(DiagramRenderBlocked):
        build(sample, preset)


@pytest.mark.parametrize('field,value', [('page_width_pt', 5000), ('page_height_pt', 5000),
                                        ('stroke', 0.00001)])
def test_unrepresentable_page_and_stroke_are_blocked(sample, field, value):
    preset = Preset()
    if field == 'stroke':
        preset.spacing.border_width_pt = value
    else:
        setattr(preset, field, value)
    with pytest.raises(DiagramRenderBlocked):
        build(sample, preset)


def test_current_writer_line_spacing_range_is_checked(sample):
    preset = Preset(page_width_pt=4032, page_height_pt=4032)
    preset.spacing.title_height = 3000
    preset.font_roles.title_pt = 1600
    with pytest.raises(DiagramRenderBlocked):
        build(sample, preset, title='A')


@pytest.mark.parametrize('field', ['title', 'eyebrow', 'subtitle', 'footnote'])
def test_common_text_measurement_overflow_returns_slot_issue(sample, field):
    metrics = FontMetrics.load_default()
    for face in (metrics.regular, metrics.bold):
        face.widths[ord('Z')] = 10**308 * face.upem
    with pytest.raises(DiagramRenderBlocked) as caught:
        build_diagram_render_plan(
            sample['diagram'], Preset(), metrics,
            evidence=[Evidence.model_validate(e) for e in sample['evidence']],
            **({'title': '정상 제목'} | {field: 'ZZ'}),
        )
    assert any(i.code == 'invalid_geometry' and i.target_id.endswith(':' + field)
               for i in caught.value.issues)


@pytest.mark.parametrize('field,text,code', [
    ('title', '제목\n둘째 줄\n셋째 줄', 'height_overflow'),
    ('eyebrow', '분류\n둘째 줄', 'height_overflow'),
    ('subtitle', '부제\n둘째 줄', 'height_overflow'),
    ('footnote', '각주\n둘째 줄\n셋째 줄', 'height_overflow'),
    ('title', '연속  공백', 'unsupported_text'),
    ('title', '제목\t탭', 'unsupported_text'),
    ('footnote', '없는 문자😀', 'unsupported_text'),
])
def test_common_text_failure_never_returns_a_partial_page(sample, field, text, code):
    with pytest.raises(DiagramRenderBlocked) as caught:
        build(sample, **{field: text})
    assert any(i.code == code and i.target_id.endswith(':' + field) for i in caught.value.issues)


def test_no_title_or_forged_input_is_rejected(sample):
    with pytest.raises(ValueError):
        build(sample, title='')
    sample['diagram']['computed'] = True
    with pytest.raises(ValueError):
        build(sample)


def test_layout_failure_propagates_and_preset_is_revalidated(sample):
    sample['diagram']['edges'][0]['to_node_id'] = 'pilot'
    with pytest.raises(DiagramRenderBlocked):
        build(sample)
    preset = Preset()
    preset.spacing.__dict__['line_spacing'] = float('nan')
    with pytest.raises(ValueError):
        build(sample, preset)


def _named(prs):
    return {s.name: s for s in prs.slides[0].shapes}


def _description(shape):
    return json.loads(shape._element.xpath('.//p:cNvPr')[0].get('descr'))


def _assert_bounds(shape, bounds):
    actual = (shape.left, shape.top, shape.width, shape.height)
    expected = (bounds.x, bounds.y, bounds.w, bounds.h)
    assert actual == tuple(round(v * 12700) for v in expected)


@pytest.mark.parametrize('relation', ['flow', 'reference'])
@pytest.mark.parametrize('stroke', [0.75, 2.0])
def test_pptx_nodes_connectors_arrows_and_evidence_are_editable(sample, tmp_path, relation, stroke):
    if relation == 'reference':
        sample['diagram']['edges'][0].update(relation='reference', from_node_id='review', to_node_id='intake')
    preset = Preset()
    preset.spacing.border_width_pt = stroke
    plan = build(sample, preset)
    out = tmp_path / 'diagram.pptx'
    write_pptx(plan, out)
    prs = Presentation(out)
    page = plan.slides[0].diagram
    shapes = _named(prs)
    prefix = page.layout.diagram_id
    for node in page.layout.nodes:
        shape = shapes[f'{prefix}:node:{node.node.id}']
        _assert_bounds(shape, node.bounds)
        assert _description(shape)['node'] == node.node.model_dump(mode='json')
    for edge in page.layout.edges:
        shape = shapes[f'{prefix}:edge:{edge.edge.id}']
        assert shape.begin_x == round(edge.start.point.x * 12700)
        assert shape.end_x == round(edge.end.point.x * 12700)
        assert shape.begin_y == shape.end_y == round(edge.start.point.y * 12700)
        assert _description(shape)['edge'] == edge.edge.model_dump(mode='json')
        dash = shape._element.xpath('.//a:custDash/a:ds')
        assert shape._element.spPr.get_or_add_ln().get('cap') == 'flat'
        assert [(int(d.get('d')), int(d.get('sp'))) for d in dash] == [
            (round(edge.dash_pattern_pt[i] / page.layout.border_width_pt * 100000),
             round(edge.dash_pattern_pt[i + 1] / page.layout.border_width_pt * 100000))
            for i in range(0, len(edge.dash_pattern_pt), 2)
        ]
        assert [int(d.get(key)) / 100000 * shape.line.width.pt for d in dash for key in ('d', 'sp')] == pytest.approx(edge.dash_pattern_pt)
        arrow_name = f'{prefix}:edge:{edge.edge.id}:arrow'
        assert (arrow_name in shapes) == bool(edge.arrow_points)
        if edge.arrow_points:
            arrow = shapes[arrow_name]
            assert _description(arrow)['edge'] == edge.edge.model_dump(mode='json')
            points = arrow._element.xpath('.//a:path/a:moveTo/a:pt | .//a:path/a:lnTo/a:pt')
            assert [(int(p.get('x')) + arrow.left, int(p.get('y')) + arrow.top) for p in points] == [
                (round(p.x * 12700), round(p.y * 12700)) for p in edge.arrow_points
            ]
    assert not prs.slides[0]._element.xpath('.//p:pic')
    body = shapes[f'{prefix}:node:intake:text:1']
    body.text_frame.paragraphs[0].runs[0].text = '수동 편집 확인'
    prs.save(out)
    assert _named(Presentation(out))[body.name].text == '수동 편집 확인'


def test_pptx_exact_lines_font_and_spacing_include_empty_lines(sample, tmp_path):
    sample['diagram']['nodes'][0]['content'] = '접수 <조건>\n\nA & B'
    plan = build(sample)
    page = plan.slides[0].diagram
    out = tmp_path / 'lines.pptx'
    write_pptx(plan, out)
    shapes = _named(Presentation(out))
    expected = []
    prefix = page.layout.diagram_id
    for n in page.layout.nodes:
        for i, text in enumerate(n.texts):
            expected.append((f'{prefix}:node:{n.node.id}:text:{i}', text.lines, text.font_pt,
                             text.line_height_pt, text.bounds))
    for e in page.layout.edges:
        t = e.label
        expected.append((f'{prefix}:edge:{e.edge.id}:label', t.lines, t.font_pt, t.line_height_pt, t.bounds))
    for h in page.headers:
        p = h.paras[0]
        expected.append((h.name, p.lines, p.font_pt, p.font_pt * plan.style.line_spacing, h))
    for name, lines, font, line_height, bounds in expected:
        shape = shapes[name]
        _assert_bounds(shape, bounds)
        tf = shape.text_frame
        assert tf.word_wrap is False and tf.auto_size == MSO_AUTO_SIZE.NONE
        assert tf.margin_left == tf.margin_right == tf.margin_top == tf.margin_bottom == 0
        assert [p.text for p in tf.paragraphs] == lines
        for p in tf.paragraphs:
            assert abs(p.line_spacing.pt - line_height) < 0.001
            assert p.space_before.pt == p.space_after.pt == 0
            for run in p.runs:
                assert run.font.size.pt == font
                assert run.font.name == 'Noto Sans KR'
                assert run._r.get_or_add_rPr().find(qn('a:ea')).get('typeface') == 'Noto Sans KR'


def test_shared_browser_fixture_matches_real_calculation():
    fixture = json.loads((Path(__file__).parent / 'fixtures/q3b-render.json').read_text('utf-8'))
    result = build(fixture['input'], **fixture['headers'])
    assert result.model_dump(mode='json') == fixture['render_plan']
