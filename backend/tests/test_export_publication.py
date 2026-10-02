"""Separate processes exercise publication through the web and CLI entry points."""

import errno
import hashlib
import io
import multiprocessing
import os
import threading
from concurrent.futures import ThreadPoolExecutor
from contextlib import redirect_stderr, redirect_stdout
from pathlib import Path

import pytest
from fastapi.testclient import TestClient
from pptx import Presentation

from slidecaptain.__main__ import main
from slidecaptain.export import exporter
from slidecaptain.models.deck import Bullet, BulletBoxSlots, Chapter, Deck, DeckMeta, Slide, Structure
from slidecaptain.models.quality import QualityReport
from slidecaptain.server.app import create_app
from slidecaptain.storage.file_store import FileProjectStore


# Windows spawn imports the full application before signaling test events.
# Allow for that startup without changing the production lock timeout.
_PROCESS_TIMEOUT = 60 if os.name == "nt" else 15


def _deck(text="Synthetic original"):
    return Deck(
        meta=DeckMeta(title="동시 내보내기"),
        structure=Structure(chapters=[Chapter(id="c1", topic="Overview", template="bullet_box")]),
        slides=[Slide(chapter_id="c1", slots=BulletBoxSlots(
            bullets=[Bullet(text=text)], conclusion="Draft conclusion",
        ))],
    )


def _prepare(tmp_path):
    root = tmp_path / "프로젝트 폴더"
    store = FileProjectStore(root)
    store.create_project("web", title="Web")
    store.save_deck("web", _deck("Synthetic web input"))
    cli_path = tmp_path / "CLI 자료.json"
    cli_path.write_text(_deck("Synthetic CLI input").model_dump_json(), encoding="utf-8")
    return root, store.exports_dir("web"), cli_path


def _export_worker(root, out_dir, cli_path, kind, result, *, chosen=None, release=None,
                   rendered=None, after_record=None, fail_pptx=False):
    """No provider is installed. Events stop a real export at a known boundary."""
    selected = []
    real_next = exporter._next_version_path
    real_write = exporter.write_pptx
    real_link = os.link

    def choose(*args):
        path = real_next(*args)
        selected.append(str(path))
        if chosen is not None:
            chosen.set()
        if release is not None and after_record is None:
            assert release.wait(_PROCESS_TIMEOUT), "publication release timed out"
        return path

    def write(*args):
        real_write(*args)
        if rendered is not None:
            rendered.set()

    def link(source, destination, **kwargs):
        if Path(destination).suffix == ".pptx" and fail_pptx:
            raise OSError(errno.EIO, "simulated second publication failure")
        real_link(source, destination, **kwargs)
        if str(destination).endswith(".quality.json") and after_record is not None:
            after_record.set()
            assert release.wait(_PROCESS_TIMEOUT), "record release timed out"

    exporter._next_version_path = choose
    exporter.write_pptx = write
    os.link = link
    try:
        if kind == "web":
            with TestClient(create_app(FileProjectStore(root)),
                            headers={"X-Requested-With": "SlideCaptain"}) as client:
                response = client.post("/api/projects/web/export")
                result.put({"kind": kind, "status": response.status_code, "body": response.json()})
        else:
            stdout, stderr = io.StringIO(), io.StringIO()
            with redirect_stdout(stdout), redirect_stderr(stderr):
                code = main(["export", str(cli_path), "--out", str(out_dir)])
            result.put({"kind": kind, "status": code, "paths": selected,
                        "stdout": stdout.getvalue(), "stderr": stderr.getvalue()})
    except BaseException as exc:
        result.put({"kind": kind, "error": repr(exc)})
        raise


def _finish(process):
    process.join(_PROCESS_TIMEOUT)
    if process.is_alive():
        process.kill()
        process.join(5)
        pytest.fail("export process did not exit")
    assert process.exitcode == 0


def _cleanup(processes):
    for process in processes:
        if process.is_alive():
            process.kill()
        process.join(5)
        assert not process.is_alive()


def _report(path):
    report = QualityReport.model_validate_json(path.with_suffix(".quality.json").read_text(encoding="utf-8"))
    assert report.artifact_sha256 == hashlib.sha256(path.read_bytes()).hexdigest()
    assert len(Presentation(str(path)).slides) == report.slide_count == 1
    assert report.final_export_allowed is False
    return report


def _artifact_bytes(out_dir):
    return {p.name: p.read_bytes() for p in out_dir.iterdir()
            if p.suffix == ".pptx" or p.name.endswith(".quality.json")}


@pytest.mark.parametrize("first_kind", ["web", "cli"])
def test_web_and_cli_processes_serialize_version_selection(tmp_path, first_kind):
    root, out_dir, cli_path = _prepare(tmp_path)
    first_path = exporter.export_deck_data(_deck(), out_dir)
    original = _artifact_bytes(out_dir)
    input_bytes = {p: p.read_bytes() for p in [root / "web" / "deck.json", cli_path]}
    ctx = multiprocessing.get_context("spawn")
    result = ctx.Queue()
    chosen, release, second_chosen, second_rendered = [ctx.Event() for _ in range(4)]
    second_kind = "cli" if first_kind == "web" else "web"
    first = ctx.Process(target=_export_worker, args=(root, out_dir, cli_path, first_kind, result),
                        kwargs={"chosen": chosen, "release": release})
    second = ctx.Process(target=_export_worker, args=(root, out_dir, cli_path, second_kind, result),
                         kwargs={"chosen": second_chosen, "rendered": second_rendered})
    processes = []
    try:
        first.start()
        processes.append(first)
        assert chosen.wait(_PROCESS_TIMEOUT)
        second.start()
        processes.append(second)
        assert second_rendered.wait(_PROCESS_TIMEOUT)
        assert not second_chosen.wait(0.3), "another process selected a version while publication was locked"
        release.set()
        rows = [result.get(timeout=_PROCESS_TIMEOUT) for _ in range(2)]
        for process in processes:
            _finish(process)
    finally:
        release.set()
        _cleanup(processes)
        result.close()
        result.join_thread()
    web = next(row for row in rows if row["kind"] == "web")
    cli = next(row for row in rows if row["kind"] == "cli")
    assert web["status"] == 200, web
    assert cli["status"] == 0, cli
    web_path, cli_output = Path(web["body"]["path"]), Path(cli["paths"][-1])
    assert {web_path.name, cli_output.name} == {"동시 내보내기_v002.pptx", "동시 내보내기_v003.pptx"}
    assert web["body"]["quality"] == _report(web_path).model_dump(mode="json")
    assert _report(web_path).input_fingerprint != _report(cli_output).input_fingerprint
    for path, marker in [(web_path, "Synthetic web input"), (cli_output, "Synthetic CLI input")]:
        texts = [shape.text for slide in Presentation(str(path)).slides for shape in slide.shapes if shape.has_text_frame]
        assert marker in texts
    assert str(cli_output) in cli["stdout"]
    assert len(list(out_dir.glob("*.pptx"))) == 3
    assert all((out_dir / name).read_bytes() == data for name, data in original.items())
    assert all(p.read_bytes() == data for p, data in input_bytes.items())
    assert first_path.exists()
    assert list(out_dir.glob(".slidecaptain-export-*/"))  # private staging is intentionally retained


def test_failed_second_publication_in_web_process_preserves_version_for_cli(tmp_path):
    root, out_dir, cli_path = _prepare(tmp_path)
    exporter.export_deck_data(_deck(), out_dir)
    original = _artifact_bytes(out_dir)
    ctx = multiprocessing.get_context("spawn")
    result = ctx.Queue()
    process = ctx.Process(target=_export_worker, args=(root, out_dir, cli_path, "web", result),
                          kwargs={"fail_pptx": True})
    process.start()
    try:
        row = result.get(timeout=_PROCESS_TIMEOUT)
        _finish(process)
        assert row["status"] == 422, row
        assert isinstance(row["body"]["detail"], str)
        orphan = out_dir / "동시 내보내기_v002.quality.json"
        orphan_bytes = orphan.read_bytes()
        assert not (out_dir / "동시 내보내기_v002.pptx").exists()
        assert main(["export", str(cli_path), "--out", str(out_dir)]) == 0
        following = out_dir / "동시 내보내기_v003.pptx"
        _report(following)
        assert orphan.read_bytes() == orphan_bytes
        assert all((out_dir / name).read_bytes() == data for name, data in original.items())
    finally:
        _cleanup([process])
        result.close()
        result.join_thread()


@pytest.mark.parametrize("kill_after_record", [False, True])
def test_killed_export_process_releases_lock_without_reusing_orphan(tmp_path, kill_after_record):
    root, out_dir, cli_path = _prepare(tmp_path)
    exporter.export_deck_data(_deck(), out_dir)
    original = _artifact_bytes(out_dir)
    ctx = multiprocessing.get_context("spawn")
    result = ctx.Queue()
    reached, release = ctx.Event(), ctx.Event()
    kwargs = {"release": release, "after_record" if kill_after_record else "chosen": reached}
    process = ctx.Process(target=_export_worker, args=(root, out_dir, cli_path, "web", result), kwargs=kwargs)
    process.start()
    try:
        assert reached.wait(_PROCESS_TIMEOUT)
        process.kill()
        process.join(5)
        assert not process.is_alive()
        orphan = out_dir / "동시 내보내기_v002.quality.json"
        assert orphan.exists() is kill_after_record
        assert not (out_dir / "동시 내보내기_v002.pptx").exists()
        assert main(["export", str(cli_path), "--out", str(out_dir)]) == 0
        version = 3 if kill_after_record else 2
        _report(out_dir / f"동시 내보내기_v{version:03d}.pptx")
        assert all((out_dir / name).read_bytes() == data for name, data in original.items())
    finally:
        _cleanup([process])
        result.close()
        result.join_thread()


def test_timeout_is_retryable_in_web_and_cli_and_next_export_succeeds(tmp_path, monkeypatch, capsys):
    from slidecaptain.export.locking import export_directory_lock

    root, out_dir, cli_path = _prepare(tmp_path)
    monkeypatch.setattr(exporter, "export_directory_lock", lambda path: export_directory_lock(path, timeout=0.05))
    with export_directory_lock(out_dir):
        with TestClient(create_app(FileProjectStore(root)), headers={"X-Requested-With": "SlideCaptain"}) as client:
            response = client.post("/api/projects/web/export")
        assert response.status_code == 409
        assert "다시" in response.json()["detail"]
        assert main(["export", str(cli_path), "--out", str(out_dir)]) == 1
        assert "다시" in capsys.readouterr().err
        assert _artifact_bytes(out_dir) == {}
    assert main(["export", str(cli_path), "--out", str(out_dir)]) == 0
    _report(out_dir / "동시 내보내기_v001.pptx")
    assert list(out_dir.glob(".slidecaptain-export-*/"))  # private staging is intentionally retained


@pytest.mark.parametrize("destination_suffix", [".pptx", ".quality.json"])
def test_destination_appearing_after_scan_is_never_overwritten(tmp_path, monkeypatch, destination_suffix):
    real_next = exporter._next_version_path
    occupied = []

    def choose(*args):
        path = real_next(*args)
        occupied.append(path.with_suffix(destination_suffix))
        occupied[-1].write_bytes(b"External user's existing file")
        return path

    monkeypatch.setattr(exporter, "_next_version_path", choose)
    with pytest.raises(FileExistsError):
        exporter.export_deck_data(_deck(), tmp_path)
    assert occupied[0].read_bytes() == b"External user's existing file"


def test_high_orphan_version_and_similar_title_do_not_reuse_or_skip_versions(tmp_path):
    (tmp_path / "동시 내보내기_v010.quality.json").write_bytes(b"reserved version")
    (tmp_path / "동시 내보내기_v020_v900.pptx").write_bytes(b"another title")
    path = exporter.export_deck_data(_deck(), tmp_path)
    assert path.name == "동시 내보내기_v011.pptx"
    assert (tmp_path / "동시 내보내기_v010.quality.json").read_bytes() == b"reserved version"
    _report(path)


def test_lock_file_is_stable_other_directories_are_independent_and_exceptions_release(tmp_path):
    from slidecaptain.export.locking import ExportBusyError, export_directory_lock

    other = tmp_path / "other"
    other.mkdir()
    with pytest.raises(RuntimeError, match="synthetic failure"):
        with export_directory_lock(tmp_path):
            lock_stat = (tmp_path / ".slidecaptain-export.lock").stat()
            with export_directory_lock(other, timeout=0):
                pass
            with pytest.raises(ExportBusyError):
                with export_directory_lock(tmp_path, timeout=0):
                    pytest.fail("same directory was not locked")
            raise RuntimeError("synthetic failure")
    with export_directory_lock(tmp_path, timeout=0):
        assert (tmp_path / ".slidecaptain-export.lock").stat().st_ino == lock_stat.st_ino
    assert (tmp_path / ".slidecaptain-export.lock").read_bytes() == b""


def test_directory_alias_shares_publication_lock(tmp_path):
    from slidecaptain.export.locking import ExportBusyError, export_directory_lock

    actual = tmp_path / "actual"
    actual.mkdir()
    alias = tmp_path / "alias"
    try:
        alias.symlink_to(actual, target_is_directory=True)
    except OSError:
        pytest.skip("directory symlinks unavailable on this device")
    with export_directory_lock(actual):
        with pytest.raises(ExportBusyError):
            with export_directory_lock(alias, timeout=0):
                pytest.fail("directory alias bypassed the lock")


def test_lock_system_error_is_not_misreported_as_contention(tmp_path, monkeypatch):
    from slidecaptain.export import locking

    def unsupported(fd):
        raise OSError(errno.ENOSYS, "unsupported locking")

    monkeypatch.setattr(locking, "_try_lock", unsupported)
    with pytest.raises(OSError) as error:
        with locking.export_directory_lock(tmp_path):
            pytest.fail("failed lock entered publication")
    assert error.value.errno == errno.ENOSYS
    assert not isinstance(error.value, locking.ExportBusyError)


def test_direct_export_threads_share_lock_and_render_outside_it(tmp_path, monkeypatch):
    barrier = threading.Barrier(4)
    real_write = exporter.write_pptx

    def write(*args):
        real_write(*args)
        barrier.wait(timeout=10)

    monkeypatch.setattr(exporter, "write_pptx", write)
    with ThreadPoolExecutor(max_workers=4) as pool:
        futures = [pool.submit(exporter.export_deck_data, _deck(f"Synthetic input {i}"), tmp_path)
                   for i in range(4)]
        paths = [future.result(timeout=15) for future in futures]
    assert {p.name for p in paths} == {f"동시 내보내기_v{i:03d}.pptx" for i in range(1, 5)}
    assert len({_report(path).input_fingerprint for path in paths}) == 4


def test_staging_uses_output_filesystem_and_unsupported_publication_fails_closed(tmp_path, monkeypatch):
    real_write = exporter.write_pptx

    def write(plan, destination):
        assert os.fstat(destination.fileno()).st_dev == tmp_path.stat().st_dev
        assert list(tmp_path.glob(".slidecaptain-export-*"))
        real_write(plan, destination)

    def unsupported(source, destination, **kwargs):
        raise OSError(errno.ENOTSUP, "hard links unavailable")

    monkeypatch.setattr(exporter, "write_pptx", write)
    with monkeypatch.context() as scoped:
        scoped.setattr(exporter.os, "link", unsupported)
        with pytest.raises(OSError) as error:
            exporter.export_deck_data(_deck(), tmp_path)
        assert error.value.errno == errno.ENOTSUP
    assert _artifact_bytes(tmp_path) == {}
    assert list(tmp_path.glob(".slidecaptain-export-*/"))  # never delete through mutable names
    _report(exporter.export_deck_data(_deck(), tmp_path))


def test_lock_open_failure_never_enters_publication_and_is_not_a_timeout(tmp_path, monkeypatch):
    from slidecaptain.export.locking import ExportBusyError, export_directory_lock

    real_open = os.open

    def denied(path, *args, **kwargs):
        if Path(path).name == ".slidecaptain-export.lock":
            raise PermissionError(errno.EACCES, "cannot open lock file")
        return real_open(path, *args, **kwargs)

    monkeypatch.setattr(os, "open", denied)
    with pytest.raises(PermissionError) as error:
        with export_directory_lock(tmp_path):
            pytest.fail("inaccessible lock was ignored")
    assert not isinstance(error.value, ExportBusyError)
    assert list(tmp_path.iterdir()) == []
