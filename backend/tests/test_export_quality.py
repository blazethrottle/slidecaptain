"""Exported artifacts, API results and CLI checks share the Q2d source verdict."""

import hashlib
import json
from pathlib import Path

import pytest
from pptx import Presentation

from slidecaptain.__main__ import main
from slidecaptain.export.exporter import export_deck_data
from slidecaptain.layout.engine import build_render_plan
from slidecaptain.metrics.font_metrics import FontMetrics
from slidecaptain.models.deck import Deck
from slidecaptain.models.preset import Preset
from slidecaptain.models.quality import QualityReport
from slidecaptain.pipeline.quality import assess_quality


FIXTURE = Path(__file__).parent / "fixtures" / "q2d-numeric-review.json"


def sample(kind="matched"):
    fixture = json.loads(FIXTURE.read_text(encoding="utf-8"))
    deck = Deck.model_validate(fixture["deck"])
    sources = fixture["sources"]
    if kind == "mismatched":
        deck.slides[0].slots.bullets[0].text = fixture["mismatched_report"]["items"][0]["actual"]
    elif kind == "stale":
        sources["synthetic.md"] += "\nChanged source"
    elif kind == "empty":
        deck.slides[0].slots.bullets[0].text = "Qualitative statement"
    return deck, sources, fixture


def saved_report(path):
    return QualityReport.model_validate_json(path.with_suffix(".quality.json").read_text(encoding="utf-8"))


def setup_project(store, kind="matched"):
    deck, sources, fixture = sample(kind)
    store.create_project("synthetic", deck.meta.title)
    for name, text in sources.items():
        store.write_source("synthetic", name, text)
    store.save_deck("synthetic", deck)
    return deck, sources, fixture


@pytest.mark.parametrize("kind,status,evaluated,failed", [
    ("matched", "passed", 1, 0), ("mismatched", "failed", 1, 1),
    ("stale", "not_run", 0, 0), ("empty", "not_run", 0, 0),
])
def test_export_records_exact_numeric_verdict_without_quality_promotion(tmp_path, kind, status, evaluated, failed):
    deck, sources, fixture = sample(kind)
    before = deck.model_dump_json()
    path = export_deck_data(deck, tmp_path, sources=sources)
    report = saved_report(path)
    assert report.gate_version == "preflight-v2"
    check = next(c for c in report.checks if c.name == "numeric_expressions")
    assert (check.status, check.evaluated, check.failed) == (status, evaluated, failed)
    assert report.numeric_review.evaluated == evaluated
    if kind in ("matched", "mismatched"):
        assert report.numeric_review.model_dump(mode="json") == fixture[f"{kind}_report"]
    else:
        assert report.numeric_review.reason == ("stale_plan" if kind == "stale" else "no_numeric_fields")
    assert report.final_export_allowed is False
    assert all(c.status == "not_run" for c in report.checks if c.name in ("evidence", "visual", "narrative"))
    assert report.status == ("needs_revision" if kind == "mismatched" else "draft")
    assert report.artifact_sha256 == hashlib.sha256(path.read_bytes()).hexdigest()
    assert len(Presentation(str(path)).slides) == report.slide_count == 1
    assert deck.model_dump_json() == before


def test_source_change_changes_quality_identity_and_preserves_earlier_artifact(tmp_path):
    deck, sources, _ = sample()
    first = export_deck_data(deck, tmp_path, sources=sources)
    before = (first.read_bytes(), first.with_suffix(".quality.json").read_bytes())
    sources["synthetic.md"] += "\nSource changed after review"
    second = export_deck_data(deck, tmp_path, sources=sources)
    assert first != second
    assert saved_report(first).input_fingerprint != saved_report(second).input_fingerprint
    assert saved_report(second).numeric_review.reason == "stale_plan"
    assert before == (first.read_bytes(), first.with_suffix(".quality.json").read_bytes())


def test_api_response_is_the_published_record_and_preserves_project_inputs(client, store):
    _, _, fixture = setup_project(store)
    before = {p: p.read_bytes() for p in store.root.rglob("*") if p.is_file()}
    response = client.post("/api/projects/synthetic/export")
    assert response.status_code == 200
    data = response.json()
    path = Path(data["path"])
    assert data["quality_path"] == str(path.with_suffix(".quality.json"))
    assert data["quality"] == json.loads(Path(data["quality_path"]).read_text(encoding="utf-8"))
    assert data["quality"]["numeric_review"] == fixture["matched_report"]
    assert {p: p.read_bytes() for p in before} == before
    lock_path = path.parent / ".slidecaptain-export.lock"
    provenance_path = path.with_suffix(".provenance.json")
    stages = list(path.parent.glob(".slidecaptain-export-*"))
    assert len(stages) == 1  # Private staging is retained to avoid path-race deletion.
    staged = {p.name: p for p in stages[0].iterdir()}
    assert set(staged) == {"deck.pptx", "deck.quality.json", "deck.provenance.json"}
    for name, published in (("deck.pptx", path), ("deck.quality.json", Path(data["quality_path"])),
                            ("deck.provenance.json", provenance_path)):
        assert staged[name].read_bytes() == published.read_bytes()
    provenance = json.loads(provenance_path.read_text(encoding="utf-8"))
    assert provenance["input_fingerprint"] == data["quality"]["input_fingerprint"]
    assert provenance["artifact_sha256"] == data["quality"]["artifact_sha256"]
    assert set(p for p in store.root.rglob("*") if p.is_file()) - set(before) == {
        path, Path(data["quality_path"]), provenance_path, lock_path, *staged.values(),
    }
    assert lock_path.read_bytes() == b""


def test_api_reassesses_after_source_edit_without_accepting_client_verdict(client, store):
    deck, sources, fixture = setup_project(store)
    assert client.post("/api/projects/synthetic/review/numbers", json=deck.model_dump(mode="json")).json() == fixture["matched_report"]
    store.write_source("synthetic", "synthetic.md", sources["synthetic.md"] + "\nChanged")
    response = client.post("/api/projects/synthetic/export", json={"numeric_review": fixture["matched_report"]})
    assert response.status_code == 200
    assert response.json()["quality"]["numeric_review"]["reason"] == "stale_plan"


def test_api_final_gate_cannot_be_bypassed_and_publishes_nothing(client, store):
    setup_project(store)
    response = client.post("/api/projects/synthetic/export?final=true")
    assert response.status_code == 422
    assert list(store.exports_dir("synthetic").iterdir()) == []


def test_cli_quality_export_and_api_use_the_same_source_snapshot(client, store, capsys):
    setup_project(store)
    deck_path = store.root / "synthetic" / "deck.json"
    assert main(["quality", str(deck_path)]) == 2
    cli_report = QualityReport.model_validate_json(capsys.readouterr().out)
    assert main(["export", str(deck_path)]) == 0
    capsys.readouterr()
    cli_path = next(store.exports_dir("synthetic").glob("*.pptx"))
    response = client.post("/api/projects/synthetic/export")
    api_report = QualityReport.model_validate(response.json()["quality"])
    exported = saved_report(cli_path)
    assert cli_report.numeric_review.status == "matched"
    assert cli_report.input_fingerprint == exported.input_fingerprint == api_report.input_fingerprint
    assert cli_report.numeric_review == exported.numeric_review == api_report.numeric_review


def test_old_quality_record_can_be_read_without_inventing_numeric_review():
    deck, _, _ = sample()
    preset = Preset()
    plan = build_render_plan(deck, preset, FontMetrics.load_default())
    legacy = assess_quality(deck, preset, plan).model_dump(mode="json")
    legacy["gate_version"] = "preflight-v1"
    legacy.pop("numeric_review", None)
    legacy["checks"] = [c for c in legacy["checks"] if c["name"] != "numeric_expressions"]
    restored = QualityReport.model_validate(legacy)
    assert restored.numeric_review is None
    assert restored.final_export_allowed is False


@pytest.mark.parametrize("kind", ["matched", "mismatched", "stale", "empty"])
def test_shared_frontend_record_equals_the_exported_backend_result(tmp_path, kind):
    fixture = json.loads((FIXTURE.parent / "q1b1-quality.json").read_text(encoding="utf-8"))
    deck, sources, _ = sample(kind)
    report = saved_report(export_deck_data(deck, tmp_path, sources=sources))
    # The shared UI fixture is the pre-publication record; the artifact hash is assigned by export.
    assert report.model_dump(mode="json", exclude={"artifact_sha256"}) == {
        key: value for key, value in fixture[kind].items() if key != "artifact_sha256"
    }


def test_partially_linked_numeric_fields_are_counted_as_revision(tmp_path):
    deck, sources, _ = sample()
    deck.slides[0].subtitle = "Unlinked 99%"
    report = saved_report(export_deck_data(deck, tmp_path, sources=sources))
    check = next(c for c in report.checks if c.name == "numeric_expressions")
    assert (check.status, check.evaluated, check.failed, check.chapter_ids) == ("failed", 2, 1, ["c1"])
    assert (report.numeric_review.matched, report.numeric_review.unresolved) == (1, 1)
    assert report.status == "needs_revision"


def test_reopened_and_restored_decks_get_new_matching_records(client, store):
    deck, _, _ = setup_project(store)
    first = client.post("/api/projects/synthetic/export").json()
    store.snapshot_now("synthetic")
    snapshot = store.list_snapshots("synthetic")[-1].id
    deck.slides[0].slots.bullets[0].text = "-200"
    store.save_deck("synthetic", deck)
    edited = client.post("/api/projects/synthetic/export").json()
    assert edited["quality"]["numeric_review"]["status"] == "needs_review"
    assert client.post(f"/api/projects/synthetic/snapshots/{snapshot}/restore").status_code == 200
    restored = client.post("/api/projects/synthetic/export").json()
    assert restored["quality"]["input_fingerprint"] == first["quality"]["input_fingerprint"]
    assert restored["quality"]["numeric_review"]["status"] == "matched"
    assert len({result["path"] for result in (first, edited, restored)}) == 3
    for result in (first, edited, restored):
        assert saved_report(Path(result["path"])).model_dump(mode="json") == result["quality"]


@pytest.mark.parametrize("encoding", ["utf-8-sig", "cp949"])
def test_cli_and_api_decode_sources_identically_and_ignore_hidden_files(client, store, capsys, encoding):
    _, sources, _ = setup_project(store)
    project = store.root / "synthetic"
    (project / "sources" / "synthetic.md").write_bytes(sources["synthetic.md"].encode(encoding))
    (project / "sources" / ".tmp-ignored").write_bytes(b"\xff")
    assert main(["quality", str(project / "deck.json")]) == 2
    cli_report = json.loads(capsys.readouterr().out)
    api_report = client.post("/api/projects/synthetic/export").json()["quality"]
    assert cli_report["numeric_review"] == api_report["numeric_review"]
    assert cli_report["numeric_review"]["status"] == "matched"


def test_unreadable_source_never_publishes_an_export(client, store, capsys):
    setup_project(store)
    first = client.post("/api/projects/synthetic/export").json()
    # Include retained staging bytes and directory entries, preserving the same
    # no-new-publication and no-change-on-unreadable-source safety contract.
    exports = store.exports_dir("synthetic")
    before = {p: p.read_bytes() for p in exports.rglob("*") if p.is_file()}
    before_entries = set(exports.rglob("*"))
    (store.root / "synthetic/sources/synthetic.md").write_bytes(b"\xff")
    assert client.post("/api/projects/synthetic/export").status_code == 422
    deck_path = store.root / "synthetic/deck.json"
    for command in ("quality", "export"):
        assert main([command, str(deck_path)]) == 1
        captured = capsys.readouterr()
        assert captured.out == "" and captured.err
    assert {p: p.read_bytes() for p in exports.rglob("*") if p.is_file()} == before
    assert set(exports.rglob("*")) == before_entries
    assert Path(first["path"]).exists()


def test_missing_sources_remain_unreviewed_without_blocking_a_draft(tmp_path):
    deck, _, _ = sample()
    report = saved_report(export_deck_data(deck, tmp_path))
    assert report.numeric_review.status == "not_run"
    assert report.numeric_review.evaluated == 0
    assert report.numeric_review.unresolved == 1
    assert report.final_export_allowed is False
