"""The review kit must be reproducible, editable, and isolated from old runs."""

import hashlib
import importlib.util
import json
from pathlib import Path

from pptx import Presentation


def test_review_kit_exports_all_pages_and_keeps_previous_run(tmp_path):
    script = Path(__file__).resolve().parents[2] / "scripts/prepare_windows_review.py"
    spec = importlib.util.spec_from_file_location("windows_review_kit", script)
    kit = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(kit)
    first = kit.prepare(tmp_path)
    manifest = json.loads((first / "manifest.json").read_text(encoding="utf-8"))
    assert manifest["ai_calls"] == 0 and manifest["human_status"] == "not_run"
    assert len(manifest["artifacts"]) == 6  # 2026-10-18 C3: 유효숫자 15자리 값 가로 막대(bar15) 추가
    assert sum(len(a["pages"]) for a in manifest["artifacts"]) == 16
    checklist = manifest["targeted_review"]
    assert checklist["version"] == "c3-t4-v1"
    assert len(checklist["checks"]) == 6
    assert len({check["id"] for check in checklist["checks"]}) == 6
    artifacts = {item["id"]: item for item in manifest["artifacts"]}
    for check in checklist["checks"]:
        assert check["title"] and check["instruction"] and check["expected"]
        assert "verdict" not in check  # 준비가 사람의 판정을 만들지 않는다.
        if check["artifact_id"] is not None:
            assert check["page"] in {p["page"] for p in artifacts[check["artifact_id"]]["pages"]}
        else:
            assert check["id"] == "status-icon" and check["page"] is None
    originals = {}
    for item in manifest["artifacts"]:
        path = first / item["file"]
        originals[path] = path.read_bytes()
        assert hashlib.sha256(originals[path]).hexdigest() == item["sha256"]
        pptx = Presentation(path)
        assert len(pptx.slides) == len(item["pages"])
        if item["id"] in ("bar", "column"):
            charts = [s.chart for s in pptx.slides[0].shapes if s.has_chart]
            assert len(charts) == 1
            assert charts[0].series[0].values == (1200.0, 1000.0)
        if item["id"] == "bar15":
            charts = [s.chart for s in pptx.slides[0].shapes if s.has_chart]
            assert charts[0].series[0].values == (123456789.123456, 1000.0)
        quality = json.loads(path.with_suffix(".quality.json").read_text(encoding="utf-8"))
        assert quality["final_export_allowed"] is False
    saved_feedback = first / "feedback.json"
    saved_feedback.write_text('{"note":"keep my feedback"}', encoding="utf-8")
    second = kit.prepare(tmp_path)
    assert second != first and second.exists()
    assert saved_feedback.read_text(encoding="utf-8") == '{"note":"keep my feedback"}'
    assert all(p.read_bytes() == data for p, data in originals.items())
    form = (first / "index.html").read_text(encoding="utf-8")
    assert "__REVIEW_MANIFEST__" not in form
    assert manifest["artifacts"][0]["sha256"] in form
