"""Sourced chart opt-in and exact-text spans preserve original editable data."""

import hashlib
import json
import io
import zipfile
from xml.etree import ElementTree
from pathlib import Path

import pytest
from pptx import Presentation

from slidecaptain.models.deck import Deck
from slidecaptain.models.preset import Preset
from slidecaptain.metrics.font_metrics import FontMetrics
from slidecaptain.layout.engine import build_render_plan
from slidecaptain.export.pptx_writer import write_pptx
from slidecaptain.pipeline.story import _text_hash, story_fingerprint


def chart_sample(a_value='1,200'):
    # a_value는 원문과 근거의 A팀 값을 함께 바꾼다. 원문이 바뀌므로 근거의 원문 해시도 다시 계산한다
    raw=(Path(__file__).parent/'fixtures/q2d-numeric-review.json').read_text(encoding='utf-8')
    fixture=json.loads(raw.replace('1,200',a_value))
    for evidence in fixture['deck']['structure']['story_plan']['evidence']:
        evidence['source_revision']=_text_hash(fixture['sources'][evidence['source_id']])
    deck=Deck.model_validate(fixture['deck'])
    payload=deck.model_dump(mode='json')
    payload['structure']['chapters'][0]['template']='table'
    payload['slides'][0]['slots']={'template':'table','columns':['팀','순매출'],
                                  'rows':[['A',a_value.replace(',','')+'원'],['B','1000원']],'footnote':'같은 기간과 분모'}
    deck=Deck.model_validate(payload)
    deck.structure.story_plan.input_fingerprint=story_fingerprint(deck.structure.story_plan,deck.meta,deck.structure.chapters,fixture['sources'])
    return deck,fixture['sources']


def test_chart_preserves_source_table_and_writes_native_editable_chart(tmp_path):
    from slidecaptain.models.expression import ChartSpec
    deck,sources=chart_sample()
    original=deck.slides[0].slots.model_dump()
    comparison=deck.structure.story_plan.comparisons[0]
    deck.slides[0].chart=ChartSpec(rule_version='comparison-chart-v1',comparison_id=comparison.id,kind='column')
    plan=build_render_plan(deck,Preset(),FontMetrics.load_default(),sources=sources)
    chart=next(frame.chart for frame in plan.slides[0].frames if frame.chart is not None)
    assert {point.evidence_id for point in chart.points}=={comparison.left_evidence_id,comparison.right_evidence_id}
    assert deck.slides[0].slots.model_dump()==original
    output=tmp_path/'chart.pptx'
    write_pptx(plan,output)
    reopened=Presentation(output)
    charts=[shape.chart for shape in reopened.slides[0].shapes if shape.has_chart]
    assert len(charts)==1 and charts[0].series[0].values==(1200.0,1000.0)


def test_partial_bold_span_preserves_exact_text_in_plan_and_pptx(tmp_path):
    from slidecaptain.models.expression import TextSpan
    fixture=json.loads((Path(__file__).parent/'fixtures/q2d-numeric-review.json').read_text())
    deck=Deck.model_validate(fixture['deck'])
    text=deck.slides[0].slots.bullets[0].text
    deck.slides[0].text_spans=[TextSpan(slot='bullets',index=0,start=0,end=2,role='bold',text_sha256=hashlib.sha256(text.encode()).hexdigest())]
    plan=build_render_plan(deck,Preset(),FontMetrics.load_default())
    para=next(frame for frame in plan.slides[0].frames if frame.name.endswith(':bullets')).paras[0]
    assert ''.join(run.text for run in para.runs)==text
    assert para.runs[0].bold is True and para.runs[-1].bold is False
    output=tmp_path/'span.pptx'
    write_pptx(plan,output)
    prs=Presentation(output)
    shape=next(shape for shape in prs.slides[0].shapes if shape.name.endswith(':bullets'))
    assert shape.text==text
    assert shape.text_frame.paragraphs[0].runs[0].font.bold is True


def test_editable_chart_workbook_preserves_formula_and_url_labels_as_plain_text(tmp_path):
    from slidecaptain.models.expression import ChartSpec
    deck,sources=chart_sample()
    deck.slides[0].chart=ChartSpec(rule_version='comparison-chart-v1',comparison_id=deck.structure.story_plan.comparisons[0].id,kind='bar')
    plan=build_render_plan(deck,Preset(),FontMetrics.load_default(),sources=sources)
    chart=next(f.chart for f in plan.slides[0].frames if f.chart)
    chart.points[0].label='=1+1'
    chart.points[1].label='https://example.com/'
    chart.definition='=1+2'
    output=tmp_path/'literal.pptx'
    write_pptx(plan,output)
    with zipfile.ZipFile(output) as package:
        embedded=next(n for n in package.namelist() if n.startswith('ppt/embeddings/'))
        with zipfile.ZipFile(io.BytesIO(package.read(embedded))) as workbook:
            sheet=ElementTree.fromstring(workbook.read('xl/worksheets/sheet1.xml'))
            ns={'s':'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
            assert not sheet.findall('.//s:f',ns)
            assert not sheet.findall('.//s:hyperlink',ns)
            strings=ElementTree.fromstring(workbook.read('xl/sharedStrings.xml'))
            texts=[node.text for node in strings.findall('.//s:t',ns)]
            assert '=1+1' in texts and 'https://example.com/' in texts
            assert any(text.startswith('=1+2') for text in texts)


@pytest.mark.parametrize('kind',['bar','column'])
def test_chart_values_are_source_bound_and_geometries_stay_in_frame(kind):
    from slidecaptain.models.expression import ChartSpec
    deck,sources=chart_sample()
    deck.slides[0].chart=ChartSpec(rule_version='comparison-chart-v1',comparison_id=deck.structure.story_plan.comparisons[0].id,kind=kind)
    frame=next(frame for frame in build_render_plan(deck,Preset(),FontMetrics.load_default(),sources=sources).slides[0].frames if frame.chart)
    chart=frame.chart
    assert [p.value for p in chart.points]==['1200','1000']
    assert chart.axis_minimum=='0'
    assert chart.axis_maximum=='1320'  # 종전 '1320.0'이 미리보기 축에 그대로 보였다 (2026-10-16)
    assert '단위 원' in chart.conditions and '분모 해당 없음' in chart.conditions
    for point in chart.points:
        assert 0<=point.x<=point.x+point.w<=frame.w
        assert 0<=point.y<=point.y+point.h<=frame.h


@pytest.mark.parametrize('maximum,expected',[('1320.00','1320'),('13.20','13.2'),('1.10','1.1'),('1','1'),('5','5'),
                                             ('1100000000000.00','1100000000000'),('135802468.04','135802468.04')])
def test_axis_maximum_text_keeps_only_the_decimals_the_maximum_needs(maximum,expected):
    # 2026-10-16 회사 PC 실측: 축 최댓값 문자열 '1320.0'이 미리보기 축에 그대로 보였다. 2026-10-20 회사 PC 이월표 (16)
    # 사용자 결정 A로 눈금은 0과 최댓값 둘뿐이라, 최댓값 자체의 자릿수만 남긴다(종전에는 다섯 눈금의 자릿수를 남겼다)
    from decimal import Decimal
    from slidecaptain.layout.expression import axis_maximum_text
    assert axis_maximum_text(Decimal(maximum))==expected


def _shown(format_code,value):
    """PowerPoint가 이 서식 기호로 값을 보일 글자. '0', '0.00' 꼴과, 0만 '0'으로 보이는 '0.00;-0.00;0' 꼴만 받는다.
    PowerPoint는 float 값을 유효숫자 15자리 십진수로 읽은 뒤 사사오입한다(2026-10-20 회사 PC 실측: 1.005, 2.675,
    123456789.125가 '0.00'에서 1.01, 2.68, 123456789.13. Python의 '.2f'는 1.005를 1.00으로 보인다)."""
    import re
    from decimal import ROUND_HALF_UP, Decimal
    match=re.fullmatch(r'(0(?:\.(0+))?)(?:;-\1;(0))?',format_code)
    assert match, f'미리보기와 같은 글자를 보장하지 못하는 서식: {format_code}'
    if match.group(3) and float(value)==0:
        return match.group(3)
    places=len(match.group(2) or '')
    return format(Decimal(format(float(value),'.15g')).quantize(Decimal(1).scaleb(-places),rounding=ROUND_HALF_UP),'f')


def _in_schema_order(element,order):
    tags=[child.tag.split('}')[1] for child in element]
    known=[tag for tag in tags if tag in order]
    return known==sorted(known,key=order.index)


_DLBL_ORDER=['idx','delete','layout','tx','numFmt','spPr','txPr','dLblPos','showLegendKey','showVal','showCatName',
             'showSerName','showPercent','showBubbleSize','separator','extLst']
_DLBLS_ORDER=['dLbl','delete','numFmt','spPr','txPr','dLblPos','showLegendKey','showVal','showCatName','showSerName',
              'showPercent','showBubbleSize','separator','showLeaderLines','leaderLines','extLst']


def _chart_labels(path,count):
    """점마다 실제로 적용될 라벨 설정. 점별(c:dLbl), 계열(c:ser/c:dLbls), 차트(c:barChart/c:dLbls) 순으로 앞선 값이 이긴다."""
    ns={'c':'http://schemas.openxmlformats.org/drawingml/2006/chart',
        'a':'http://schemas.openxmlformats.org/drawingml/2006/main'}
    with zipfile.ZipFile(path) as package:
        root=ElementTree.fromstring(package.read(next(n for n in package.namelist() if n.startswith('ppt/charts/chart'))))
    plot=root.find('.//c:plotArea/c:barChart/c:dLbls',ns)
    series=root.find('.//c:ser/c:dLbls',ns)
    points={int(d.find('c:idx',ns).get('val')):d for d in (series.findall('c:dLbl',ns) if series is not None else [])}

    def effective(index,path,attr):
        for level in (points.get(index),series,plot):
            node=level.find(path,ns) if level is not None else None
            if node is not None:
                return node.get(attr)
        return None

    labels=[{'format':effective(i,'c:numFmt','formatCode'),'show':effective(i,'c:showVal','val'),
             'position':effective(i,'c:dLblPos','val'),'size':effective(i,'c:txPr/a:p/a:pPr/a:defRPr','sz')}
            for i in range(count)]
    # ElementTree의 find는 순서를 보지 않으므로, PowerPoint가 복구 오류를 내는 스키마 순서 위반을 따로 본다
    for container in (plot,series):
        if container is not None:
            assert _in_schema_order(container,_DLBLS_ORDER)
    for dlbl in points.values():
        assert _in_schema_order(dlbl,_DLBL_ORDER)
    # PowerPoint가 빠진 설정을 위 단위에서 물려받는다는 가정에 기대지 않는다. 계열이나 점 단위 라벨을 만들었다면
    # 그 요소가 서식, 값 표시, 위치, 글꼴을 스스로 갖고 있어야 한다 (2026-10-16 후속 리뷰)
    for own in ([series] if series is not None else [])+list(points.values()):
        for path in ('c:numFmt','c:showVal','c:dLblPos','c:txPr/a:p/a:pPr/a:defRPr'):
            assert own.find(path,ns) is not None, path
    axis=root.find('.//c:valAx/c:numFmt',ns).get('formatCode')
    return labels,axis


@pytest.mark.parametrize('kind',['bar','column'])
@pytest.mark.parametrize('a_value,values,axis_text',[('1,200',None,'1320'),('1,200',(('1200','1200'),('1000.5','1000.5')),'1320'),
                                                     ('1,234',None,'1357.4'),('123456789.123456',None,'135802468.04'),
                                                     ('1,200',(('1.005','1.01'),('2.675','2.68')),'1320')])
def test_pptx_chart_labels_and_axis_show_the_same_text_as_the_preview(tmp_path,kind,a_value,values,axis_text):
    # 2026-10-16 회사 PC 실측(PowerPoint 16.0 build 20430): '0.##############' 서식 때문에 값과 눈금이 '1200.',
    # '1320.'처럼 끝에 마침표가 붙었다. 미리보기는 표시 글자(display)를 그리고 배치도 그 글자 폭으로 계산한다.
    # 소수 3자리 이상인 값은 소수 2자리로 보인다(2026-10-20 회사 PC 이월표 (18) 사용자 결정 A). values는 (원문, 표시) 쌍이다
    from slidecaptain.models.expression import ChartSpec
    deck,sources=chart_sample(a_value)
    deck.slides[0].chart=ChartSpec(rule_version='comparison-chart-v1',comparison_id=deck.structure.story_plan.comparisons[0].id,kind=kind)
    plan=build_render_plan(deck,Preset(),FontMetrics.load_default(),sources=sources)
    chart=next(f.chart for f in plan.slides[0].frames if f.chart)
    assert chart.axis_maximum==axis_text
    for point,(value,display) in zip(chart.points,values or ()):
        point.value,point.display=value,display
    output=tmp_path/'chart.pptx'
    write_pptx(plan,output)
    labels,axis=_chart_labels(output,len(chart.points))
    for point,label in zip(chart.points,labels):
        # 계열이나 점에 라벨 설정을 새로 만들면 값 표시, 위치, 글꼴이 차트 단위 설정을 덮으므로 함께 본다
        assert label['show']=='1'
        assert label['position']=='outEnd'
        assert label['size']==str(int(chart.font_pt*100))
        assert _shown(label['format'],point.value)==point.display
    assert _shown(axis,chart.axis_maximum)==chart.axis_maximum
    assert _shown(axis,chart.axis_minimum)=='0'  # 미리보기는 0 눈금을 '0'으로 그린다


@pytest.mark.parametrize('a_value,display',[('1,200','1200'),('1,234.5','1234.5'),('1,234.56','1234.56'),
                                            ('123456789.123456','123456789.12'),('2.999','3.00'),('1.005','1.01'),
                                            ('0.004','0.00')])
def test_chart_shows_at_most_two_decimals_and_keeps_the_source_value(tmp_path,a_value,display):
    # 2026-10-20 회사 PC 이월표 (18) 사용자 결정 A: 소수 3자리 이상인 차트 값은 소수 셋째 자리에서 사사오입해 2자리로
    # 보인다. 2자리 이하는 그대로 보이고 0을 채우지 않는다. 원문 값은 렌더 계획과 PPTX 내장 데이터에 그대로 남는다
    from slidecaptain.models.expression import ChartSpec
    deck,sources=chart_sample(a_value)
    deck.slides[0].chart=ChartSpec(rule_version='comparison-chart-v1',comparison_id=deck.structure.story_plan.comparisons[0].id,kind='bar')
    plan=build_render_plan(deck,Preset(),FontMetrics.load_default(),sources=sources)
    chart=next(f.chart for f in plan.slides[0].frames if f.chart)
    source=a_value.replace(',','')
    assert [(p.value,p.display) for p in chart.points]==[(source,display),('1000','1000')]
    output=tmp_path/'chart.pptx'
    write_pptx(plan,output)
    reopened=next(shape.chart for shape in Presentation(output).slides[0].shapes if shape.has_chart)
    assert reopened.series[0].values==(float(source),1000.0)
    descr=next(shape for shape in Presentation(output).slides[0].shapes if shape.has_chart)._element.nvGraphicFramePr.cNvPr.get('descr')
    assert json.loads(descr)['points'][0]['value']==source


def test_preview_value_width_is_measured_on_the_display_text():
    # 값 글자 자리는 표시 글자 폭으로 잡는다. 원문 폭으로 잡으면 15자리 값에서 그림 영역이 쓸데없이 좁아진다
    from slidecaptain.models.expression import ChartSpec
    deck,sources=chart_sample('123456789.123456')
    deck.slides[0].chart=ChartSpec(rule_version='comparison-chart-v1',comparison_id=deck.structure.story_plan.comparisons[0].id,kind='bar')
    metrics=FontMetrics.load_default()
    frame=next(f for f in build_render_plan(deck,Preset(),metrics,sources=sources).slides[0].frames if f.chart)
    chart=frame.chart
    width=metrics.face(False).width_pt('123456789.12',chart.font_pt)
    assert chart.plot_w==pytest.approx(frame.w-chart.plot_x-width-chart.font_pt*2)


@pytest.mark.parametrize('kind',['bar','column'])
def test_pptx_value_axis_has_no_gridlines_and_only_zero_and_maximum_ticks(tmp_path,kind):
    # 2026-10-20 회사 PC 이월표 (16) 사용자 결정 A: PowerPoint가 주 눈금선과 중간 눈금 글자(264~1056)를 그렸고
    # 미리보기는 0과 최댓값만 그렸다. bar15에서는 값 글자가 최댓값 눈금선과 겹쳤다. PowerPoint를 미리보기에 맞춘다
    from slidecaptain.models.expression import ChartSpec
    ns={'c':'http://schemas.openxmlformats.org/drawingml/2006/chart'}
    deck,sources=chart_sample()
    deck.slides[0].chart=ChartSpec(rule_version='comparison-chart-v1',comparison_id=deck.structure.story_plan.comparisons[0].id,kind=kind)
    plan=build_render_plan(deck,Preset(),FontMetrics.load_default(),sources=sources)
    chart=next(f.chart for f in plan.slides[0].frames if f.chart)
    output=tmp_path/'chart.pptx'
    write_pptx(plan,output)
    with zipfile.ZipFile(output) as package:
        root=ElementTree.fromstring(package.read(next(n for n in package.namelist() if n.startswith('ppt/charts/chart'))))
    assert root.findall('.//c:valAx/c:majorGridlines',ns)==[] and root.findall('.//c:valAx/c:minorGridlines',ns)==[]
    assert root.findall('.//c:catAx/c:majorGridlines',ns)==[]
    assert [float(node.get('val')) for node in root.findall('.//c:valAx/c:majorUnit',ns)]==[float(chart.axis_maximum)]


def test_column_chart_leaves_room_below_the_axis_for_powerpoint_category_labels():
    # 2026-10-16 회사 PC 실측(PowerPoint 16.0 build 20430, 12pt): 범주 이름 글자는 x축 아래 15.5~25.5pt에 그려진다.
    # 축 아래 공간이 24pt(글자 크기의 2배)라 글자 아랫부분이 차트 틀에서 잘렸다. 종전 축 서식('0.###')에서는
    # PowerPoint가 그림 영역을 계획보다 9pt 위로 줄여 그려 가려져 있었다(그 대신 미리보기와 어긋났다)
    from slidecaptain.models.expression import ChartSpec
    deck,sources=chart_sample()
    deck.slides[0].chart=ChartSpec(rule_version='comparison-chart-v1',comparison_id=deck.structure.story_plan.comparisons[0].id,kind='column')
    frame=next(f for f in build_render_plan(deck,Preset(),FontMetrics.load_default(),sources=sources).slides[0].frames if f.chart)
    chart=frame.chart
    assert frame.h-(chart.plot_y+chart.plot_h)>=25.5/12*chart.font_pt+2


def test_bar_chart_leaves_room_below_the_plot_for_powerpoint_value_axis_labels():
    # 2026-10-20 회사 PC 실측(PowerPoint 16.0 build 20430, 12pt): 값 축 글자 칸은 약 28pt다. 그림 영역 아래 공간이
    # 24pt(글자 크기의 2배)면 PowerPoint가 글자 칸을 차트 틀 안에 넣으려고 그림 영역을 계획보다 4.2pt 위에 그렸다.
    # 같은 PPTX의 높이만 바꿔 보면 28pt에서 0.16pt 위, 30pt에서 계획과 일치했다
    from slidecaptain.models.expression import ChartSpec
    deck,sources=chart_sample()
    deck.slides[0].chart=ChartSpec(rule_version='comparison-chart-v1',comparison_id=deck.structure.story_plan.comparisons[0].id,kind='bar')
    frame=next(f for f in build_render_plan(deck,Preset(),FontMetrics.load_default(),sources=sources).slides[0].frames if f.chart)
    chart=frame.chart
    assert frame.h-(chart.plot_y+chart.plot_h)>=28.2/12*chart.font_pt


def test_column_chart_leaves_room_for_the_axis_maximum_the_preview_draws():
    # 리뷰 발견(2026-10-16): 1234와 1000의 축 최댓값 '1357.40'이 값 글자보다 넓어, 미리보기가 축 글자를
    # 프레임 왼쪽 밖(-4.7pt)에서 그리기 시작해 앞자리가 잘렸다. 미리보기는 축 최댓값을 PowerPoint 실측처럼
    # plot_x-font*1.125에 끝 맞춤한다(2026-10-20 회사 PC, 종전 font/2)
    from slidecaptain.models.expression import ChartSpec
    deck,sources=chart_sample('1,234')
    deck.slides[0].chart=ChartSpec(rule_version='comparison-chart-v1',comparison_id=deck.structure.story_plan.comparisons[0].id,kind='column')
    metrics=FontMetrics.load_default()
    chart=next(f.chart for f in build_render_plan(deck,Preset(),metrics,sources=sources).slides[0].frames if f.chart)
    assert chart.axis_maximum=='1357.4'
    assert chart.plot_x-chart.font_pt*1.125-metrics.face(False).width_pt(chart.axis_maximum,chart.font_pt)>=0


@pytest.mark.parametrize('damage',['missing_sources','stale_source','incompatible','forged_comparison','value_token'])
def test_chart_never_falls_back_to_original_table_for_unverifiable_values(damage):
    from slidecaptain.models.expression import ChartSpec
    deck,sources=chart_sample()
    deck.slides[0].chart=ChartSpec(rule_version='comparison-chart-v1',comparison_id=deck.structure.story_plan.comparisons[0].id,kind='bar')
    if damage=='missing_sources':
        sources=None
    elif damage=='stale_source':
        sources={name:text+'\nchanged' for name,text in sources.items()}
    elif damage=='incompatible':
        deck.structure.story_plan.evidence[1].metric_basis.definition='다른 지표'
    elif damage=='forged_comparison':
        deck.slides[0].chart.comparison_id='unknown'
    else:
        deck.structure.story_plan.evidence[0].value='200'  # substring of 1200 is not a fact
    with pytest.raises(ValueError):
        build_render_plan(deck,Preset(),FontMetrics.load_default(),sources=sources)


@pytest.mark.parametrize('changes',[{'kind':'pie'},{'values':[1,2]},{'x':0},{'approved':True}])
def test_chart_does_not_accept_client_values_geometry_or_approval(changes):
    from slidecaptain.models.expression import ChartSpec
    with pytest.raises(ValueError):
        ChartSpec.model_validate({'rule_version':'comparison-chart-v1','comparison_id':'cp1','kind':'bar',**changes})


@pytest.mark.parametrize('changes',[{'start':True},{'end':0},{'end':100000},{'index':1},
                                   {'text_sha256':'0'*64},{'role':'arbitrary'},{'slot':'rows'},
                                   {'index':None},{'start':2,'end':2}])
def test_spans_reject_invalid_offsets_stale_text_or_wrong_scope(changes):
    fixture=json.loads((Path(__file__).parent/'fixtures/q2d-numeric-review.json').read_text())
    payload=fixture['deck']
    text=payload['slides'][0]['slots']['bullets'][0]['text']
    payload['slides'][0]['text_spans']=[{'slot':'bullets','index':0,'start':0,'end':2,'role':'bold',
                                       'text_sha256':hashlib.sha256(text.encode()).hexdigest(),**changes}]
    with pytest.raises(ValueError):
        Deck.model_validate(payload)


def test_overlapping_spans_and_edits_cannot_rebind_old_offsets():
    from slidecaptain.models.expression import TextSpan
    fixture=json.loads((Path(__file__).parent/'fixtures/q2d-numeric-review.json').read_text())
    deck=Deck.model_validate(fixture['deck'])
    text=deck.slides[0].slots.bullets[0].text
    span={'slot':'bullets','index':0,'start':0,'end':3,'role':'bold','text_sha256':hashlib.sha256(text.encode()).hexdigest()}
    deck.slides[0].text_spans=[TextSpan(**span),TextSpan(**{**span,'start':2,'end':4,'role':'accent'})]
    with pytest.raises(ValueError,match='겹칠'):
        build_render_plan(deck,Preset(),FontMetrics.load_default())
    deck.slides[0].text_spans=deck.slides[0].text_spans[:1]
    deck.slides[0].slots.bullets[0].text='다른 문장'
    with pytest.raises(ValueError,match='기준 문장'):
        build_render_plan(deck,Preset(),FontMetrics.load_default())


def test_legacy_deck_and_numeric_fingerprint_are_identical_without_opt_in():
    from slidecaptain.pipeline.numeric_review import assess_numeric_review
    fixture=json.loads((Path(__file__).parent/'fixtures/q2d-numeric-review.json').read_text())
    deck=Deck.model_validate(fixture['deck'])
    assert deck.model_dump(mode='json')==fixture['deck']
    assert assess_numeric_review(deck,fixture['sources']).model_dump(mode='json')==fixture['matched_report']
    plan=build_render_plan(deck,Preset(),FontMetrics.load_default())
    for frame in plan.model_dump(mode='json')['slides'][0]['frames']:
        assert 'chart' not in frame
        assert all('runs' not in para and 'line_runs' not in para for para in frame['paras'])


def test_bold_face_width_changes_wrapping_and_is_not_replaced_by_regular_width():
    from slidecaptain.layout.expression import _styled_lines
    from slidecaptain.models.expression import TextRun
    class Face:
        def __init__(self,width):self.width=width
        def width_pt(self,text,font):return len(text)*self.width
    class Metrics:
        def face(self,bold):return Face(10 if bold else 1)
    plain=_styled_lines([TextRun(text='AAAA BBBB',bold=False,color='202020')],8,12,Metrics(),1)
    styled=_styled_lines([TextRun(text='AAAA',bold=True,color='202020'),TextRun(text=' BBBB',bold=False,color='202020')],8,12,Metrics(),1)
    assert len(styled)>len(plain)
    assert ''.join(run.text for line in styled for run in line)=='AAAA BBBB'


def _chart_with_height(kind,height_in_fonts):
    # 차트 틀 높이(조건 문장 칸을 뺀 높이)를 글자 크기의 배수로 맞춘다. 조건 칸 높이는 넉넉한 틀에서 한 번 잰다
    from slidecaptain.models.expression import ChartSpec
    from slidecaptain.models.render import Frame
    from slidecaptain.layout.expression import build_comparison_chart
    deck,sources=chart_sample()
    deck.slides[0].chart=ChartSpec(rule_version='comparison-chart-v1',comparison_id=deck.structure.story_plan.comparisons[0].id,kind=kind)
    preset,metrics=Preset(),FontMetrics.load_default()
    _chart,caption=build_comparison_chart(deck,deck.slides[0],Frame(name='c',x=0,y=0,w=600,h=400),preset,metrics,sources)
    font=preset.font_roles.body_pt
    frame=Frame(name='c',x=0,y=0,w=600,h=caption.h+font*height_in_fonts)
    return lambda: build_comparison_chart(deck,deck.slides[0],frame,preset,metrics,sources)


@pytest.mark.parametrize('kind,lower',[('column',8.5),('bar',8)])
def test_narrow_chart_height_fails_with_the_single_space_message(kind,lower):
    # 회사 PC 이월표 (13): 세로 막대는 축 아래 공간을 글자 크기의 2.5배로 늘린 뒤(1b106aa) 높이 8~8.5배 구간에서
    # 둘째 문구("차트 축과 값을 배치할...")로 실패했다. 좁은 높이는 종류와 관계없이 첫 문구 하나로 안내한다
    for height in (lower-0.25,lower-0.01):
        with pytest.raises(ValueError,match='차트와 필수 조건을 배치할 공간이 부족합니다'):
            _chart_with_height(kind,height)()
    chart,_caption=_chart_with_height(kind,lower)()
    assert chart.chart.plot_h>=chart.chart.font_pt*4


@pytest.mark.parametrize('a_value',['1234567890.123456','12345678901.23456','999999999999.123456'])
def test_chart_rejects_values_beyond_fifteen_significant_digits(a_value):
    # 회사 PC 이월표 (7): PowerPoint는 유효숫자 15자리까지 보이고 미리보기는 원문 자릿수를 그대로 보인다.
    # 16자리 값은 float 왕복을 통과해도 두 화면의 숫자가 달라지므로 그리지 않는다 (2026-10-18 사용자 결정 Q3 A).
    # float로 왕복하지 않는 18자리 값도 같은 안내를 받는다(C3 리뷰 C3R-04: 검사 순서 때문에 예전 문구로 거절됐다)
    from slidecaptain.models.expression import ChartSpec
    deck,sources=chart_sample(a_value)
    deck.slides[0].chart=ChartSpec(rule_version='comparison-chart-v1',comparison_id=deck.structure.story_plan.comparisons[0].id,kind='bar')
    with pytest.raises(ValueError,match='15자리'):
        build_render_plan(deck,Preset(),FontMetrics.load_default(),sources=sources)


def test_fifteen_digit_value_draws_and_its_axis_maximum_has_at_most_two_decimals():
    # 15자리 값은 그린다. 축 최댓값은 값의 1.1배를 소수 2자리로 올린 값이다(135802468.0358016 -> 135802468.04).
    # 눈금이 0과 최댓값뿐이라(회사 PC 이월표 (16)) 최댓값 자체만 15자리 안이면 된다
    from decimal import Decimal
    from slidecaptain.models.expression import ChartSpec
    deck,sources=chart_sample('123456789.123456')
    deck.slides[0].chart=ChartSpec(rule_version='comparison-chart-v1',comparison_id=deck.structure.story_plan.comparisons[0].id,kind='column')
    chart=next(f.chart for f in build_render_plan(deck,Preset(),FontMetrics.load_default(),sources=sources).slides[0].frames if f.chart)
    assert chart.points[0].value=='123456789.123456'
    assert chart.axis_maximum=='135802468.04'


@pytest.mark.parametrize('kind,crosses',[('bar','max'),('column','autoZero')])
def test_bar_chart_value_axis_crosses_at_the_bottom_like_the_preview(tmp_path,kind,crosses):
    # 회사 PC 이월표 (5): 가로 막대는 첫 범주를 위에 두려고 범주 순서를 뒤집는다(maxMin). 값 축은 첫 범주 쪽에서
    # 교차하므로 PowerPoint가 값 축을 위에 그렸다. 미리보기는 값 축 글자를 아래에 그린다.
    # PowerPoint 표시는 회사 PC 렌더로 확인한다. 이 시험은 XML 설정만 붙잡는다
    from slidecaptain.models.expression import ChartSpec
    ns={'c':'http://schemas.openxmlformats.org/drawingml/2006/chart'}
    deck,sources=chart_sample()
    deck.slides[0].chart=ChartSpec(rule_version='comparison-chart-v1',comparison_id=deck.structure.story_plan.comparisons[0].id,kind=kind)
    output=tmp_path/'chart.pptx'
    write_pptx(build_render_plan(deck,Preset(),FontMetrics.load_default(),sources=sources),output)
    with zipfile.ZipFile(output) as package:
        root=ElementTree.fromstring(package.read(next(n for n in package.namelist() if n.startswith('ppt/charts/chart'))))
    assert [node.get('val') for node in root.findall('.//c:valAx/c:crosses',ns)]==[crosses]
    assert [node.get('val') for node in root.findall('.//c:catAx/c:crosses',ns)]==['autoZero']
    orientation=root.find('.//c:catAx/c:scaling/c:orientation',ns)
    assert (orientation is not None and orientation.get('val')=='maxMin')==(kind=='bar')


@pytest.mark.parametrize('a_value',['500000000.00001','550000000.00001','999999999999.99','1000000000000'])
def test_axis_maximum_covers_the_value_within_two_decimals_and_fifteen_digits(a_value):
    # C3 리뷰 C3R-03은 최댓값이 15자리일 때 눈금(최댓값의 k/5)이 16자리가 되는 문제였다. 2026-10-20부터 눈금은 0과
    # 최댓값뿐이고(회사 PC 이월표 (16)) 최댓값은 소수 2자리로 올린다((18)). 값이 1조 이하라 15자리를 넘지 않는다
    from decimal import Decimal
    from slidecaptain.models.expression import ChartSpec
    deck,sources=chart_sample(a_value)
    deck.slides[0].chart=ChartSpec(rule_version='comparison-chart-v1',comparison_id=deck.structure.story_plan.comparisons[0].id,kind='bar')
    chart=next(f.chart for f in build_render_plan(deck,Preset(),FontMetrics.load_default(),sources=sources).slides[0].frames if f.chart)
    maximum=Decimal(chart.axis_maximum)
    assert maximum>=Decimal(a_value)*Decimal('1.1')
    assert -maximum.as_tuple().exponent<=2
    assert len(maximum.normalize().as_tuple().digits)<=15


def test_render_plan_routes_answer_422_with_the_chart_reason(tmp_path):
    # C3 리뷰 C3R-01: 차트를 그릴 수 없는 이유가 500 "요청이 실패했습니다"로만 보였다. 다시 시도해도 같으므로 문구를 보인다
    from fastapi.testclient import TestClient
    from slidecaptain.models.expression import ChartSpec
    from slidecaptain.server.app import create_app
    from slidecaptain.storage.file_store import FileProjectStore
    deck,sources=chart_sample('1234567890.123456')
    deck.slides[0].chart=ChartSpec(rule_version='comparison-chart-v1',comparison_id=deck.structure.story_plan.comparisons[0].id,kind='bar')
    store=FileProjectStore(tmp_path/'projects')
    store.create_project('p1')
    for filename,text in sources.items():
        store.write_source('p1',filename,text)
    store.save_deck('p1',deck,snapshot=False)
    with TestClient(create_app(store),headers={'X-Requested-With':'SlideCaptain'}) as client:
        for response in (client.get('/api/projects/p1/render-plan'),
                         client.post('/api/projects/p1/render-plan',json=deck.model_dump(mode='json'))):
            assert response.status_code==422
            assert response.json()=={'detail':'차트 값은 유효숫자 15자리까지 지원합니다. 원문 수치의 자릿수를 줄여 주세요.',
                                      'code':'chart_unavailable'}
