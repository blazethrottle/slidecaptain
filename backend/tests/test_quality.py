"""Q1a contracts using synthetic inputs; no AI or target-app quality claims."""

import hashlib
import json
from pathlib import Path

import pytest
from fastapi.testclient import TestClient
from pptx import Presentation
from pydantic import ValidationError

from slidecaptain.__main__ import main
from slidecaptain.export import exporter
from slidecaptain.layout.engine import build_render_plan
from slidecaptain.metrics.font_metrics import FontMetrics
from slidecaptain.models.deck import (
    Bullet, BulletBoxSlots, Chapter, CoverSlots, Deck, DeckMeta, DividerSlots, Slide, Structure,
)
from slidecaptain.models.preset import Preset
from slidecaptain.models.quality import QualityReport
from slidecaptain.pipeline.prompts import (
    QUALITY_RULES, _draft_context, build_chapter_prompt, build_structure_prompt,
)
from slidecaptain.pipeline.quality import (
    QualityExportBlocked, assess_quality, require_export_allowed,
)
from slidecaptain.server.app import create_app
from slidecaptain.storage.file_store import FileProjectStore


def _deck(chapters=1, generated=None, text="Synthetic evidence"):
    generated = chapters if generated is None else generated
    return Deck(
        meta=DeckMeta(title="Quality"),
        structure=Structure(chapters=[
            Chapter(id=f"c{i}", topic=f"Topic {i}", template="bullet_box")
            for i in range(chapters)
        ]),
        slides=[
            Slide(chapter_id=f"c{i}", slots=BulletBoxSlots(
                bullets=[Bullet(text=f"{text} {i}")], conclusion="Draft conclusion",
            ))
            for i in range(generated)
        ],
    )


def _inputs(deck):
    preset = Preset()
    return preset, build_render_plan(deck, preset, FontMetrics.load_default())


def _checks(report):
    return {check.name: check for check in report.checks}


def _sidecar(path):
    return QualityReport.model_validate_json(path.with_suffix(".quality.json").read_text(encoding="utf-8"))


def test_complete_preflight_is_still_an_unreviewed_draft():
    deck = _deck()
    report = assess_quality(deck, *_inputs(deck))
    assert report.status == "draft"
    assert report.draft_export_allowed is True
    assert report.final_export_allowed is False
    assert report.artifact_sha256 is None
    for name in ("nonempty_output", "chapter_coverage", "layout_capacity"):
        check = _checks(report)[name]
        assert (check.status, check.evaluated, check.failed) == ("passed", 1, 0)
    for name in ("narrative", "evidence", "representation", "visual", "target_renderer"):
        check = _checks(report)[name]
        assert (check.status, check.evaluated, check.failed) == ("not_run", 0, 0)
    require_export_allowed(report)
    with pytest.raises(QualityExportBlocked):
        require_export_allowed(report, final=True)
    with pytest.raises(ValidationError):
        QualityReport.model_validate({**report.model_dump(), "final_export_allowed": True})


@pytest.mark.parametrize("chapters", [0, 2])
def test_empty_output_is_not_a_pass_and_does_not_publish_files(tmp_path, chapters):
    deck = _deck(chapters, generated=0)
    report = assess_quality(deck, *_inputs(deck))
    checks = _checks(report)
    assert report.status == "needs_revision"
    assert report.draft_export_allowed is False
    assert report.slide_count == 0
    assert (checks["nonempty_output"].evaluated, checks["nonempty_output"].failed) == (1, 1)
    coverage = checks["chapter_coverage"]
    assert (coverage.evaluated, coverage.failed) == (chapters, chapters)
    assert coverage.status == ("failed" if chapters else "not_run")
    assert checks["layout_capacity"].status == "not_run"
    assert checks["layout_capacity"].evaluated == 0
    out_dir = tmp_path / "not-created"
    with pytest.raises(QualityExportBlocked):
        exporter.export_deck_data(deck, out_dir)
    assert not out_dir.exists()


def test_missing_chapters_are_counted_and_named_without_blocking_draft_export(tmp_path):
    deck = _deck(3, generated=1)
    report = assess_quality(deck, *_inputs(deck))
    check = _checks(report)["chapter_coverage"]
    assert (check.status, check.evaluated, check.failed) == ("failed", 3, 2)
    assert check.chapter_ids == ["c1", "c2"]
    path = exporter.export_deck_data(deck, tmp_path)
    assert _sidecar(path).status == "needs_revision"
    assert len(Presentation(str(path)).slides) == 1


def test_capacity_warning_is_a_revision_not_a_final_certificate(tmp_path):
    deck = _deck(text="Capacity overflow " * 900)
    report = assess_quality(deck, *_inputs(deck))
    check = _checks(report)["layout_capacity"]
    assert (check.status, check.evaluated, check.failed) == ("failed", 1, 1)
    assert check.chapter_ids == ["c0"]
    assert report.status == "needs_revision"
    path = exporter.export_deck_data(deck, tmp_path)
    assert _sidecar(path).draft_export_allowed is True
    assert _sidecar(path).final_export_allowed is False


def test_final_rejection_happens_before_output_directory_creation(tmp_path):
    out_dir = tmp_path / "final"
    with pytest.raises(QualityExportBlocked):
        exporter.export_deck_data(_deck(), out_dir, final=True)
    assert not out_dir.exists()


@pytest.mark.parametrize("changed", ["deck", "preset", "plan"])
def test_fingerprint_is_stable_and_covers_each_input(changed):
    deck = _deck()
    preset, plan = _inputs(deck)
    before = (deck.model_dump_json(), preset.model_dump_json(), plan.model_dump_json())
    first = assess_quality(deck, preset, plan).input_fingerprint
    assert assess_quality(deck, preset, plan).input_fingerprint == first
    assert (deck.model_dump_json(), preset.model_dump_json(), plan.model_dump_json()) == before
    if changed == "deck":
        deck.meta.title = "Changed report title"
    elif changed == "preset":
        preset.font_roles.title_pt += 1
    else:
        plan.slides[0].frames[0].x += 1
    assert assess_quality(deck, preset, plan).input_fingerprint != first


def test_versioned_sidecars_match_exact_pptx_bytes_and_preserve_deck(tmp_path):
    deck = _deck()
    before = deck.model_dump_json()
    expected = assess_quality(deck, *_inputs(deck)).input_fingerprint
    paths = [exporter.export_deck_data(deck, tmp_path) for _ in range(2)]
    assert [path.name for path in paths] == ["Quality_v001.pptx", "Quality_v002.pptx"]
    for path in paths:
        report = _sidecar(path)
        assert report.input_fingerprint == expected
        assert report.artifact_sha256 == hashlib.sha256(path.read_bytes()).hexdigest()
        assert report.slide_count == len(Presentation(str(path)).slides) == 1
    assert deck.model_dump_json() == before


def test_existing_sidecar_reserves_its_version(tmp_path):
    orphan = tmp_path / "Quality_v001.quality.json"
    orphan.write_text("existing record", encoding="utf-8")
    path = exporter.export_deck_data(_deck(), tmp_path)
    assert path.name == "Quality_v002.pptx"
    assert orphan.read_text(encoding="utf-8") == "existing record"


@pytest.mark.parametrize("failed_link", [1, 2])
def test_publish_failure_does_not_overwrite_prior_artifacts(tmp_path, monkeypatch, failed_link):
    deck = _deck()
    first = exporter.export_deck_data(deck, tmp_path)
    original_pptx = first.read_bytes()
    original_record = first.with_suffix(".quality.json").read_bytes()
    real_link = exporter.os.link
    calls = 0

    def fail_link(source, destination, **kwargs):
        nonlocal calls
        calls += 1
        if calls == failed_link:
            raise OSError("simulated publication failure")
        return real_link(source, destination, **kwargs)

    with monkeypatch.context() as scoped:
        scoped.setattr(exporter.os, "link", fail_link)
        with pytest.raises(OSError, match="simulated publication failure"):
            exporter.export_deck_data(deck, tmp_path)
    assert first.read_bytes() == original_pptx
    assert first.with_suffix(".quality.json").read_bytes() == original_record
    assert not (tmp_path / "Quality_v002.pptx").exists()
    orphan = tmp_path / "Quality_v002.quality.json"
    assert orphan.exists() is (failed_link == 2)
    saved_orphan = orphan.read_bytes() if orphan.exists() else None
    following = exporter.export_deck_data(deck, tmp_path)
    assert following.name == ("Quality_v003.pptx" if failed_link == 2 else "Quality_v002.pptx")
    if saved_orphan is not None:
        assert orphan.read_bytes() == saved_orphan
    assert _sidecar(following).artifact_sha256 == hashlib.sha256(following.read_bytes()).hexdigest()


def test_writer_failure_does_not_publish_a_record_or_pptx(tmp_path, monkeypatch):
    def fail_write(*args):
        raise OSError("simulated writer failure")

    monkeypatch.setattr(exporter, "write_pptx", fail_write)
    with pytest.raises(OSError, match="simulated writer failure"):
        exporter.export_deck_data(_deck(), tmp_path)
    assert not list(tmp_path.glob("*.pptx"))
    assert not list(tmp_path.glob("*.quality.json"))
    assert list(tmp_path.glob(".slidecaptain-export-*/"))  # private failed staging is retained


def test_quality_cli_returns_machine_readable_pending_status_without_export(tmp_path, capsys):
    path = tmp_path / "deck.json"
    path.write_text(_deck().model_dump_json(), encoding="utf-8")
    assert main(["quality", str(path)]) == 2
    captured = capsys.readouterr()
    report = QualityReport.model_validate_json(captured.out)
    assert report.status == "draft"
    assert captured.err == ""
    assert sorted(p.name for p in tmp_path.iterdir()) == ["deck.json"]


@pytest.mark.parametrize("exists", [True, False])
def test_quality_cli_input_errors_are_not_quality_verdicts(tmp_path, capsys, exists):
    path = tmp_path / "deck.json"
    if exists:
        path.write_text("{invalid", encoding="utf-8")
    assert main(["quality", str(path)]) == 1
    captured = capsys.readouterr()
    assert captured.out == ""
    assert captured.err


@pytest.mark.parametrize("final", [True, False])
def test_cli_rejects_final_or_empty_export_before_creating_output(tmp_path, capsys, final):
    path = tmp_path / "deck.json"
    path.write_text(_deck() .model_dump_json() if final else _deck(0).model_dump_json(), encoding="utf-8")
    out_dir = tmp_path / "exports"
    args = ["export", str(path), "--out", str(out_dir)]
    if final:
        args.append("--final")
    assert main(args) == 1
    captured = capsys.readouterr()
    assert captured.out == ""
    assert captured.err
    assert not out_dir.exists()


@pytest.mark.parametrize("generated", [0, 1])
def test_api_export_extends_success_and_keeps_existing_error_contract(tmp_path, generated):
    store = FileProjectStore(tmp_path / "projects")
    store.create_project("p1", title="Quality")
    store.save_deck("p1", _deck(1, generated=generated))
    with TestClient(create_app(store)) as client:
        response = client.post("/api/projects/p1/export", headers={"X-Requested-With": "SlideCaptain"})
    if not generated:
        assert response.status_code == 422
        assert set(response.json()) == {"detail"}
        assert isinstance(response.json()["detail"], str)
        assert not list(store.exports_dir("p1").glob("*.pptx"))
        assert not list(store.exports_dir("p1").glob("*.quality.json"))
    else:
        assert response.status_code == 200
        assert set(response.json()) == {"path", "quality_path", "quality"}
        path = Path(response.json()["path"])
        assert _sidecar(path).artifact_sha256 == hashlib.sha256(path.read_bytes()).hexdigest()


def test_generation_prompts_include_editorial_rules_and_other_drafts():
    deck = _deck(2)
    structure = build_structure_prompt(deck.meta, {})
    chapter = build_chapter_prompt(deck, deck.structure.chapters[0], {}, {}, today="2026-09-12")
    assert QUALITY_RULES in structure
    assert QUALITY_RULES in chapter
    assert _draft_context(deck, deck.structure.chapters[0]) in chapter
    assert "사실의 근거가 아님" in chapter


def test_draft_context_preserves_structure_order_and_excludes_current_cover_divider():
    deck = _deck(3)
    deck.slides[0].slots.bullets[0].text = "CURRENT_ONLY"
    cover = Chapter(id="cover", topic="Cover", template="cover")
    divider = Chapter(id="divider", topic="Divider", template="divider")
    deck.structure.chapters.extend([cover, divider])
    deck.slides.extend([
        Slide(chapter_id="cover", slots=CoverSlots(title="COVER_ONLY")),
        Slide(chapter_id="divider", slots=DividerSlots(section_title="DIVIDER_ONLY")),
    ])
    deck.slides.reverse()
    context = _draft_context(deck, deck.structure.chapters[0])
    assert context.index("[c1]") < context.index("[c2]")
    for excluded in ("CURRENT_ONLY", "COVER_ONLY", "DIVIDER_ONLY", "[c0]"):
        assert excluded not in context
    assert _draft_context(deck, cover) == ""
    assert _draft_context(deck, divider) == ""


def test_draft_context_is_bounded_and_marks_omissions():
    deck = _deck(12, text="Long draft " * 1000)
    context = _draft_context(deck, deck.structure.chapters[0])
    assert len(context) <= 6200
    assert "이 초안의 일부 생략됨" in context
    assert "추가 초안은 문맥 한도로 생략됨" in context
    assert "[c11]" not in context


def test_draft_context_is_empty_without_other_generated_slides():
    deck = _deck(2, generated=1)
    assert _draft_context(deck, deck.structure.chapters[0]) == ""
