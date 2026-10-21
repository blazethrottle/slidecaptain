"""Synthetic, bounded pilot through the app API; default: prepare, no AI.

Run from backend with the repository's dev dependencies installed:
  python -m scripts.quality_pilot --output /new/local/folder [--offline]
Live use additionally requires --live --consent --max-calls N.
The limit counts provider.complete attempts, not SDK turns or billing units.
D4 live check: --effort sets the requested effort, --stop-after chapter stops after the
chapter job (one call for a matrix row), --claude-cli picks the Claude Code CLI for this run.
"""
import argparse
from contextlib import ExitStack, contextmanager
from copy import deepcopy
from datetime import datetime, timezone
import hashlib
import json
import logging
import os
from pathlib import Path
import sqlite3
import tempfile
import threading
import time
import uuid

from fastapi.testclient import TestClient
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from slidecaptain.models.deck import Deck
from slidecaptain.pipeline.auth_status import ENV_CLI, LoginStatus, _bundled_cli_path, _cli_version, resolve_cli_path
from slidecaptain.pipeline.codex import CodexConnection
from slidecaptain.pipeline.connections import (
    AIConnections, AISelection, ClaudeConnection, LoginAttempt, ModelOption,
)
from slidecaptain.pipeline.provider import ProviderCallFailed, ProviderError, ProviderResponse
from slidecaptain.pipeline.story import require_current_story, story_fingerprint
from slidecaptain.pipeline.subscription import ENV_STDERR_LOG
from slidecaptain.server.app import create_app
from slidecaptain.storage.file_store import FileProjectStore
from slidecaptain.storage.job_ledger import LEDGER_NAME

PROJECT = "quality-pilot"
BODY = "pilot-action"
# 생성 작업 하나를 기다리는 상한. 등록부터 세므로 제공자 시간 제한(300초)에 연결 상태 확인과 목록 조회의 여유를 더한다.
# 300초였을 때는 도구가 먼저 멈춰 나간 호출의 결과를 기록하지 못할 수 있었다 (D4 실제 생성 확인 리뷰 LCR-16)
JOB_WAIT_SECONDS = 360
# 실제 Claude 실행 동안 빼는 세션 연결 변수의 이름 앞부분. 이 도구를 Claude Code 세션 안에서 돌리면 이 변수들이 CLI에 물려져
# 데스크톱 앱 경로와 조건이 달라진다 (리뷰 LCR-06). 사용자가 일부러 둔 설정(CLAUDE_CONFIG_DIR 등)은 데스크톱 앱처럼 남긴다
# (코드 리뷰 LCC-03). 2026-10-21 이 PC의 세션에서 본 11개를 덮는다
_SESSION_ENV_PREFIXES = (
    "CLAUDECODE", "CLAUDE_CODE_SESSION_", "CLAUDE_CODE_MESSAGING_", "CLAUDE_CODE_BRIDGE_", "CLAUDE_CODE_CHILD_SESSION",
    "CLAUDE_CODE_EXECPATH", "CLAUDE_CODE_ENTRYPOINT", "CLAUDE_PID", "CLAUDE_EFFORT",
)
_TERMINAL = {"succeeded", "failed", "cancelled", "interrupted", "remote_completion_unknown"}


def _write(path, data):
    fd, name = tempfile.mkstemp(dir=path.parent, prefix=f".{path.name}-", suffix=".tmp")
    try:
        with os.fdopen(fd, "w", encoding="utf-8", newline="\n") as stream:
            stream.write(json.dumps(data, ensure_ascii=False, indent=2) + "\n")
            stream.flush()
            os.fsync(stream.fileno())
        os.replace(name, path)
    finally:
        Path(name).unlink(missing_ok=True)


def scenario():
    fixture = Path(__file__).resolve().parents[1] / "tests/fixtures/q3b-project.json"
    sample = json.loads(fixture.read_text(encoding="utf-8"))
    data = sample["deck"]
    data["structure"]["chapters"].append({
        "id": BODY, "topic": "자동 알림 도입 전 확인할 사항", "template": "bullet_box",
        "conclusion": "효과와 비용을 확인한 뒤 도입 여부를 결정한다",
        "source_refs": list(sample["sources"]),
    })
    plan = data["structure"]["story_plan"]
    plan["claims"].append({
        "id": "action", "statement": "자동 알림의 효과와 비용을 확인한 뒤 도입 여부를 결정한다.",
        "kind": "proposal", "evidence_ids": ["e3"], "caveats": ["효과와 비용은 확인하지 않았다."],
    })
    plan["chapters"].extend([
        {"chapter_id": "cover", "role": "cover", "claim_ids": []},
        {"chapter_id": BODY, "role": "action", "claim_ids": ["action"]},
    ])
    data["slides"].append({"chapter_id": BODY, "slots": {
        "template": "bullet_box", "bullets": [{"text": "효과와 비용을 확인할 필요가 있다."}],
        "conclusion": "합성 초안: 도입 여부 미결정", "footnote": "합성 자료",
    }})
    deck = Deck.model_validate(data)
    plan = deck.structure.story_plan
    plan.input_fingerprint = story_fingerprint(plan, deck.meta, deck.structure.chapters, sample["sources"])
    require_current_story(deck, sample["sources"])
    brief = plan.brief.model_copy(update={"decision_question": "요청 처리 흐름과 자동 알림 도입 전 확인할 사항은 무엇인가?"})
    return {
        "synthetic": True, "sources": sample["sources"], "deck": deck.model_dump(mode="json"),
        "rewrite_brief": brief.model_dump(mode="json"),
        "chapter_instructions": "합성 운영 사례다. 근거에 없는 수치와 효과를 만들지 말고, 사실과 제안을 구분해 짧게 작성한다.",
        "rewrite_instructions": "기존 도식과 보호 주장/근거를 보존한다. 표지, 개선 검토, 업무 흐름 순으로 배치한다. 장은 추가하거나 삭제하지 않는다.",
    }


class OfflineProvider:
    """Fixed synthetic responses, never a real provider fallback."""

    async def complete(self, prompt, schema):
        data = scenario()
        if "chapter_order" in schema.get("properties", {}):
            plan = data["deck"]["structure"]["story_plan"]
            payload = {k: deepcopy(plan[k]) for k in (
                "claims", "answer_claim_ids", "unanswered_questions", "comparisons", "derivations",
            )}
            payload["evidence"] = []
            payload["claims"] = [claim for claim in payload["claims"] if claim["id"] == "action"]
            by_id = {c["chapter_id"]: c for c in plan["chapters"]}
            payload["chapter_order"] = ["cover", BODY, "synthetic-flow"]
            payload["chapters"] = [by_id[key] for key in ("cover", BODY)]
        else:
            payload = {
                "template": "bullet_box", "bullets": [
                    {"text": "접수 후 검토 담당자에게 요청을 전달한다."},
                    {"text": "자동 알림 도입은 제안이며 효과와 비용을 확인해야 한다."},
                ], "conclusion": "도입 전 효과와 비용을 확인한다", "footnote": "합성 자료",
            }
        return ProviderResponse(structured=payload, raw_text="synthetic offline response")


class BudgetedProvider:
    """One shared dispatch budget covers generation, retry and condense."""

    def __init__(self, upstream, *, max_calls, on_change=None, on_request=None):
        if type(max_calls) is not int or max_calls < 1:
            raise ValueError("호출 상한은 양의 정수여야 합니다.")
        self.upstream = upstream
        self.max_calls = max_calls
        self.calls = 0
        self.blocked_calls = 0
        self.records = []
        self.on_change = on_change or (lambda: None)
        self.on_request = on_request or (lambda number, prompt, schema: None)
        self._lock = threading.Lock()
        self._journal_failed = False

    async def complete(self, prompt, schema):
        with self._lock:
            if self._journal_failed:
                raise OSError("실행 기록을 보존하지 못해 추가 호출을 중단했습니다.")
            if self.calls >= self.max_calls:
                self.blocked_calls += 1
                self.on_change()
                raise ProviderCallFailed("승인된 호출 상한에 도달했습니다.")
            self.calls += 1
            record = {"number": self.calls, "status": "started", "usage": None,
                      "prompt_sha256": hashlib.sha256(prompt.encode("utf-8")).hexdigest(),
                      "schema_sha256": hashlib.sha256(json.dumps(schema, sort_keys=True).encode("utf-8")).hexdigest()}
            self.records.append(record)
            try:
                self.on_request(self.calls, prompt, schema)
                self.on_change()
            except BaseException:
                self.calls -= 1
                record["status"] = "not_sent"
                self._journal_failed = True
                raise
        try:
            response = await self.upstream.complete(prompt, schema)
        except BaseException as exc:
            record["status"] = "failed" if isinstance(exc, Exception) else "interrupted"
            usage = exc.usage if isinstance(exc, ProviderError) else None
            record["usage"] = usage.model_dump(mode="json") if usage else None
            raise
        else:
            record["status"] = "returned"
            record["usage"] = response.usage.model_dump(mode="json") if response.usage else None
            return response
        finally:
            try:
                self.on_change()
            except BaseException:
                self._journal_failed = True
                raise


class _PilotConnection:
    def __init__(self, provider, model, real=None, effort=None):
        self.bounded = provider
        self.model = model
        self.real = real
        self.effort = effort

    def status(self):
        return self.real.status() if self.real else LoginStatus(logged_in=True, auth_method="synthetic")

    def models(self):
        if self.real:
            return self.real.models()  # 실제 목록과 대조한다. 없는 모델이나 강도는 생성 임대가 호출 전에 막는다
        return [ModelOption(id=self.model, label=self.model, supported_efforts=[self.effort] if self.effort else [])]

    def login_status(self):
        return LoginAttempt(state="idle")

    def provider(self, model, *, effort=None):
        if model != self.model or effort != self.effort:
            raise ProviderCallFailed("파일럿 모델이나 추론 강도가 바뀌었습니다.")
        return self.bounded

    def close(self):
        if self.real:
            self.real.close()


class PilotStopped(Exception):
    pass


def _expect(condition, message):
    if not condition:
        raise PilotStopped(message)


def _workflow(client, output, data, report, stop_after=None):
    route = f"/api/projects/{PROJECT}"

    def request(method, suffix, *, status=200, **kwargs):
        response = client.request(method, route + suffix, **kwargs)
        _expect(response.status_code == status, f"{method} {suffix}: HTTP {response.status_code}")
        return response

    def job(kind, params, headers=None, require_result=True):
        """작업 API로 생성을 등록하고 끝날 때까지 기다린다. 동기 생성 래퍼는 D3c-7에서 지웠다. 형식 오류로 실패한 작업도
        결과(원문과 사용량)를 남기므로 결과를 돌려주고, 결과 없이 끝나면 멈춘다(require_result=False면 종결 행을 그대로 돌려준다)."""
        registered = request("POST", "/jobs", status=202, headers=headers or {}, json={
            "request_id": f"pilot-{uuid.uuid4().hex}", "kind": kind, "params": params})
        job_id = registered.json()["id"]
        report.setdefault("job_ids", []).append(job_id)  # 원장과 사용량의 요청 강도 대조에 쓴다
        deadline = time.monotonic() + JOB_WAIT_SECONDS
        while (view := request("GET", f"/jobs/{job_id}").json())["state"] not in _TERMINAL:
            _expect(time.monotonic() < deadline, f"{kind} 작업이 {JOB_WAIT_SECONDS}초 안에 끝나지 않았습니다.")
            time.sleep(0.05)
        _expect(not require_result or view["result"] is not None, f"{kind} 작업이 결과 없이 끝났습니다({view['state']}).")
        return view

    report["stage"] = "chapter"
    if stop_after == "chapter":
        # 호출 행렬의 한 칸은 장 생성 1회로 판정한다. 형식 재시도나 자동 축약이 상한에 막혀 작업이 결과 없이 끝나도 종결
        # 상태를 기록한다. 판정은 calls[0]과 진단 로그로 한다 (D4 실제 생성 확인 3.5, 리뷰 LCR-04)
        ended = job("chapter", {"chapter_id": BODY, "instructions": data["chapter_instructions"]}, require_result=False)
        checks = report["checks"]
        checks["chapter_job_state"] = ended["state"]
        checks["chapter_error_code"] = (ended.get("error") or {}).get("code")
        checks["chapter_status"] = (ended["result"] or {}).get("status")
        if ended["result"] is not None:
            _write(output / "chapter-result.json", ended["result"])
        _expect(report["provider_calls"] > 0, "호출을 보내기 전에 장 작업이 끝났습니다(호출 전 차단).")
        return
    base = request("GET", "/deck")
    chapter = job("chapter", {"chapter_id": BODY, "instructions": data["chapter_instructions"]})
    generated = chapter["result"]
    _write(output / "chapter-result.json", generated)
    _expect(generated["status"] == "ok", "장 생성 응답을 적용하지 못했습니다.")
    # 제품과 같은 반영 경로(합치기, 반영 기록, 처분)로 넣는다. 덱 PUT으로 슬롯만 바꾸면 그 경로를 재지 않는다(β 묶음 리뷰 S10)
    request("POST", f"/jobs/{chapter['id']}/apply", json={}, headers={"If-Match": base.headers["etag"]})
    _expect(request("GET", f"/jobs/{chapter['id']}").json()["candidate_status"] == "applied",
            "장 결과를 반영했지만 처분이 applied가 아닙니다.")
    before = request("GET", "/deck")
    _write(output / "before-rewrite.json", before.json())
    # The service intentionally preserves a valid draft when condensing fails.
    # That does not mean the pilot may disregard a blocked dispatch.
    _expect(report["blocked_calls"] == 0, "호출 상한으로 중단했습니다. 생성된 본문은 저장했습니다.")
    snapshots = request("GET", "/snapshots").json()

    report["stage"] = "rewrite_preview"
    rewrite = job("rewrite", {"brief": data["rewrite_brief"], "instructions": data["rewrite_instructions"]},
                  headers={"If-Match": before.headers["etag"]})
    candidate = rewrite["result"]
    _write(output / "rewrite-result.json", candidate)
    _expect(candidate["status"] == "ok", "재작성 후보를 적용하지 못했습니다.")
    # 래퍼는 만드는 동안 저장본이나 자료가 바뀌면 412나 409로 멈췄다. 작업 API에서는 후보가 낡음으로 남는다
    _expect(rewrite["candidate_status"] == "held", "재작성 후보가 만드는 동안 입력이 바뀌어 낡았습니다.")
    unchanged = request("GET", "/deck")
    checks = report["checks"]
    checks["preview_did_not_save"] = (
        unchanged.headers["etag"] == before.headers["etag"] and unchanged.json() == before.json()
        and request("GET", "/snapshots").json() == snapshots
    )
    _expect(checks["preview_did_not_save"], "미리보기 중 저장본이 바뀌었습니다.")
    report["stage"] = "rewrite_apply"
    applied = request("POST", "/story-plan/rewrite/apply", json={
        "deck": candidate["deck"], "sources_fingerprint": candidate["sources_fingerprint"],
    }, headers={"If-Match": candidate["base_etag"]}).json()
    checks["rewrite_preserved_slides"] = applied["slides"] == before.json()["slides"]
    _expect(checks["rewrite_preserved_slides"], "재작성 과정에서 기존 본문이 변경됐습니다.")
    _write(output / "applied-deck.json", applied)

    report["stage"] = "export"
    render = request("GET", "/render-plan").json()
    _write(output / "render-plan.json", render)
    exported = request("POST", "/export").json()
    _write(output / "export-result.json", exported)
    pptx = Path(exported["path"])
    quality = json.loads(Path(exported["quality_path"]).read_text(encoding="utf-8"))
    document = Presentation(pptx)
    checks["pptx_slide_count"] = len(document.slides)
    _expect(checks["pptx_slide_count"] == 3, "PPTX의 장 수가 다릅니다.")
    checks["pptx_hash_matches"] = quality["artifact_sha256"] == hashlib.sha256(pptx.read_bytes()).hexdigest()
    _expect(checks["pptx_hash_matches"], "PPTX와 품질 기록의 해시가 다릅니다.")
    shapes = {shape.name: shape for page in document.slides for shape in page.shapes}
    expected = [f"synthetic-flow:node:{node}" for node in ("intake", "review", "pilot")]
    checks["editable_diagram"] = all(name in shapes and shapes[name].shape_type != MSO_SHAPE_TYPE.PICTURE for name in expected)
    _expect(checks["editable_diagram"], "편집 가능한 도식 요소를 찾지 못했습니다.")
    paths_before = {p.name for p in pptx.parent.iterdir()}
    rejected = client.post(route + "/export?final=true")
    checks["final_export_blocked"] = rejected.status_code == 422 and paths_before == {p.name for p in pptx.parent.iterdir()}
    _expect(checks["final_export_blocked"], "미검수 제출본이 차단되지 않았습니다.")
    report["quality_status"] = quality["status"]
    report["artifacts"] = {"pptx": str(pptx), "quality": exported["quality_path"]}


def _cli_kind(cli):
    if cli is None:
        return "missing"
    if cli == _bundled_cli_path():
        return "bundled"
    return "installed" if cli.parent == Path.home() / ".local" / "bin" else "other"


def _effort_checks(output, report, effort):
    """원장 행과 사용량 기록의 요청 강도가 요청과 같은지 본다 (D4 실제 생성 확인 3.2, 3.4). 실제 적용값 대조가 아니다."""
    ids = report.get("job_ids") or []
    if not ids:
        return
    # 대조가 원장을 바꾸지 않도록 읽기 전용으로 연다. 원장의 effort는 API에 실리지 않는다 (리뷰 LCR-15)
    connection = sqlite3.connect((output / "projects" / LEDGER_NAME).as_uri() + "?mode=ro", uri=True)
    try:
        marks = ",".join("?" * len(ids))
        stored = dict(connection.execute(f"SELECT id, effort FROM jobs WHERE id IN ({marks})", ids).fetchall())
    finally:
        connection.close()
    report["checks"]["ledger_effort_matches"] = all(job_id in stored and stored[job_id] == effort for job_id in ids)
    path = output / "projects" / PROJECT / "ai-usage.jsonl"
    lines = [json.loads(line) for line in path.read_text(encoding="utf-8").splitlines()] if path.exists() else []
    # 줄이 없으면(호출 전 차단) 불일치가 아니라 판정 없음이다 (코드 리뷰 LCC-06)
    report["checks"]["usage_records"] = len(lines)
    report["checks"]["usage_effort_matches"] = (
        all(line.get("requested_effort") == effort for line in lines) if lines else None)


@contextmanager
def _pilot_environment(output, report, *, live_claude, claude_cli):
    """실행 동안만 바꾸는 환경 변수와 진단 로그. 연결 닫기를 포함해 어떤 예외로 끝나도 되돌린다 (코드 리뷰 LCC-02)."""
    with ExitStack() as cleanup:
        saved = {}

        def change(name, value):
            saved.setdefault(name, os.environ.get(name))
            if value is None:
                os.environ.pop(name, None)
            else:
                os.environ[name] = value

        def restore():
            for name, value in saved.items():
                if value is None:
                    os.environ.pop(name, None)
                else:
                    os.environ[name] = value

        cleanup.callback(restore)
        if claude_cli is not None:
            change(ENV_CLI, str(Path(claude_cli)))
        if live_claude:
            removed = sorted(name for name in os.environ if name.startswith(_SESSION_ENV_PREFIXES))
            for name in removed:
                change(name, None)
            change(ENV_STDERR_LOG, "1")  # 실패 원인을 가르려고 이 실행에서만 표준 오류를 받는다 (LCC-01)
            # 이름만 적는다. 값(세션 식별자, 토큰, 경로)은 남기지 않는다
            report["env_removed"] = removed
            report["env_present_overrides"] = sorted(
                name for name in os.environ if name.startswith(("ANTHROPIC_", "CLAUDE")))
        # 응답 관측 줄과 호출 실패 원인을 출력 폴더에 남긴다 (D4 실제 생성 확인 3.4)
        logger = logging.getLogger("slidecaptain")
        cleanup.callback(logger.setLevel, logger.level)
        diagnostic = logging.FileHandler(output / "diagnostic.log", encoding="utf-8")
        cleanup.callback(diagnostic.close)
        diagnostic.setFormatter(logging.Formatter("%(asctime)s %(levelname)s %(name)s %(message)s"))
        logger.addHandler(diagnostic)
        cleanup.callback(logger.removeHandler, diagnostic)
        logger.setLevel(logging.INFO)
        yield


def run_pilot(output, *, mode="prepare", consent=False, max_calls=None, model=None, provider=None,
              ai_provider="claude", codex_home=None, effort=None, stop_after=None, claude_cli=None):
    if mode not in {"prepare", "offline", "live"}:
        raise ValueError("지원하지 않는 실행 모드입니다.")
    if mode == "live" and (not consent or type(max_calls) is not int or max_calls < 1):
        raise ValueError("실제 호출에는 전송 동의와 명시적인 양의 호출 상한이 필요합니다.")
    limit = 5 if max_calls is None else max_calls
    if type(limit) is not int or limit < 1:
        raise ValueError("호출 상한은 양의 정수여야 합니다.")
    if ai_provider not in {"claude", "chatgpt"}:
        raise ValueError("지원하지 않는 AI 서비스입니다.")
    if stop_after not in {None, "chapter"}:
        raise ValueError("지원하지 않는 중단 단계입니다.")
    if claude_cli is not None and ai_provider != "claude":
        raise ValueError("Claude CLI는 Claude 실행에만 지정합니다.")
    model = model if model is not None else ("sonnet" if ai_provider == "claude" else "gpt-6-luna")
    # 모델은 고정 별칭이 아니라 실제 목록과 대조한다. 목록은 Claude Code 판과 계정마다 다르다 (D4b 계획 8절)
    chosen = AISelection(provider=ai_provider, model=model, effort=effort)
    output = Path(output).resolve()
    output.mkdir(parents=True, exist_ok=False)
    report = {
        "mode": mode, "status": "running", "stage": "prepare", "synthetic": True,
        "requested_provider": ai_provider, "requested_model": model, "requested_effort": effort,
        "stop_after": stop_after,
        "max_calls": limit, "provider_calls": 0, "external_calls": 0, "blocked_calls": 0,
        "call_unit": "provider.complete attempts; excludes pre-dispatch refusals; SDK turns are separate",
        "usage_basis": "app CallUsage; raw SDK missing-field normalization is not independently verified",
        "calls": [], "checks": {}, "review": {k: "not_run" for k in ("semantic", "browser", "powerpoint", "reader")},
        "started_at": datetime.now(timezone.utc).isoformat(),
    }
    with _pilot_environment(output, report, live_claude=mode == "live" and ai_provider == "claude",
                            claude_cli=claude_cli):
        return _run_pilot(output, report, mode=mode, limit=limit, model=model, provider=provider,
                          ai_provider=ai_provider, codex_home=codex_home, effort=effort, stop_after=stop_after,
                          chosen=chosen)


def _run_pilot(output, report, *, mode, limit, model, provider, ai_provider, codex_home, effort, stop_after, chosen):
    bounded = None
    real = None
    manager = None

    def persist():
        if bounded:
            report.update(provider_calls=bounded.calls, external_calls=bounded.calls if mode == "live" else 0,
                          blocked_calls=bounded.blocked_calls, calls=bounded.records)
        _write(output / "report.json", report)

    try:
        data = scenario()
        _write(output / "input.json", data)
        report["input_sha256"] = hashlib.sha256((output / "input.json").read_bytes()).hexdigest()
        store = FileProjectStore(output / "projects")
        store.create_project(PROJECT)
        for filename, text in data["sources"].items():
            store.write_source(PROJECT, filename, text)
        store.save_deck(PROJECT, Deck.model_validate(data["deck"]), snapshot=False)
        if mode == "prepare":
            report["status"] = "prepared"
            return report
        if mode == "live" and ai_provider == "claude":
            cli = resolve_cli_path()
            report["cli_kind"] = _cli_kind(cli)  # 전체 경로는 사용자 이름을 담으므로 적지 않는다
            report["cli_version"] = _cli_version(cli, 10) if cli is not None else None
        if mode == "live":
            real = (ClaudeConnection() if ai_provider == "claude" else CodexConnection(
                Path(codex_home) if codex_home is not None
                else Path.home() / "slidecaptain-projects" / ".slidecaptain-codex",
            ))
        if provider is None and real is not None:
            # 앱의 생성 임대와 같이, 강도가 없으면 강도 인자를 넘기지 않는다
            provider = real.provider(model, effort=effort) if effort is not None else real.provider(model)
            report["provider_effort"] = getattr(provider, "effort", None)  # 실제 프로바이더가 CLI에 넘길 강도 (리뷰 LCR-01)
            try:
                # 응답 모델과 대조할 목록의 해석된 모델 (리뷰 LCR-05). 목록 조회는 생성 호출이 아니다
                report["catalog_resolved_model"] = next(
                    (option.resolved_model for option in real.models() if option.id == model), None)
            except Exception as exc:  # 진단 칸이다. 목록 실패는 생성 임대가 원인 코드로 다시 막는다
                report["catalog_resolved_model"] = None
                report["catalog_error"] = type(exc).__name__
        upstream = provider or OfflineProvider()
        bounded = BudgetedProvider(upstream, max_calls=limit, on_change=persist,
                                   on_request=lambda n, prompt, schema: _write(output / f"call-{n:02d}-input.json", {
                                       "prompt": prompt, "schema": schema,
                                   }))
        connection = _PilotConnection(bounded, model, real, effort)
        manager = AIConnections(output / "ai-settings.json", initial=chosen, connections={ai_provider: connection})
        _, selection, _ = manager.selected_status()
        # pytest 밖에서도 도는 스크립트라 conftest 의 기본 Host 변경이 적용되지 않는다. 허용 Host 를 명시한다 (SV-11, PM-01)
        with TestClient(create_app(store, ai_connections=manager), base_url="http://127.0.0.1", headers={
            "X-Requested-With": "SlideCaptain", "X-AI-Consent": "SlideCaptain", "X-AI-Selection": selection,
        }) as client:
            _workflow(client, output, data, report, stop_after)
        report["status"] = "completed"
        report["stage"] = "done"
    except Exception as exc:
        report["status"] = "blocked"
        report["error"] = str(exc) if isinstance(exc, PilotStopped) else type(exc).__name__
    except BaseException:
        report["status"] = "interrupted"
        raise
    finally:
        report["finished_at"] = datetime.now(timezone.utc).isoformat()
        try:
            if mode != "prepare":
                try:
                    _effort_checks(output, report, effort)
                except Exception as exc:  # 대조 실패가 실행 결과 기록을 막지 않는다
                    report["checks"]["effort_check_error"] = type(exc).__name__
            persist()
        finally:
            if manager is not None:
                manager.close()
            elif real is not None:
                real.close()
    return report


def main(argv=None):
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output", type=Path, required=True, help="새 로컬 출력 폴더")
    modes = parser.add_mutually_exclusive_group()
    modes.add_argument("--offline", action="store_true")
    modes.add_argument("--live", action="store_true")
    parser.add_argument("--consent", action="store_true", help="합성 입력을 선택한 AI 서비스로 전송하는 데 동의")
    parser.add_argument("--max-calls", type=int)
    parser.add_argument("--provider", choices=["claude", "codex"], default="claude")
    parser.add_argument("--model", help="Claude 기본 sonnet, Codex 기본 gpt-6-luna. 다른 모델은 명시적으로 지정")
    parser.add_argument("--codex-home", type=Path, help="SlideCaptain 전용 Codex 프로필 경로 (기본: ~/slidecaptain-projects/.slidecaptain-codex)")
    parser.add_argument("--effort", help="요청할 추론 강도. 생략하면 제공자 기본값")
    parser.add_argument("--stop-after", choices=["chapter"], help="장 생성 작업이 끝나면 멈춘다(호출 행렬의 한 칸)")
    parser.add_argument("--claude-cli", type=Path, help="이 실행에서 쓸 Claude Code CLI 경로")
    args = parser.parse_args(argv)
    if args.live and (not args.consent or args.max_calls is None or args.max_calls < 1):
        parser.error("--live에는 --consent와 양의 --max-calls가 필요합니다.")
    if args.max_calls is not None and args.max_calls < 1:
        parser.error("--max-calls는 양의 정수여야 합니다.")
    try:
        report = run_pilot(args.output, mode="live" if args.live else "offline" if args.offline else "prepare",
                           consent=args.consent, max_calls=args.max_calls, model=args.model,
                           ai_provider="chatgpt" if args.provider == "codex" else "claude",
                           codex_home=args.codex_home, effort=args.effort, stop_after=args.stop_after,
                           claude_cli=args.claude_cli)
    except FileExistsError:
        parser.error("이미 존재하는 출력 폴더는 사용할 수 없습니다.")
    except ValueError:
        parser.error("AI 서비스와 모델 설정을 확인해 주세요.")
    print(json.dumps({k: report[k] for k in ("status", "stage", "provider_calls", "external_calls")}, ensure_ascii=False))
    return 0 if report["status"] in {"prepared", "completed"} else 2


if __name__ == "__main__":
    raise SystemExit(main())
